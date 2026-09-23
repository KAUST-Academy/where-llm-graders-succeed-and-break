#!/usr/bin/env python3
"""Table 15 (tab:ftpaired) marginal verdicts.

Recomputes the paired third-grader statistic d-bar (Eq. paired) and its 95% CI
to three decimals for every pooled cell of Table 15 (marks and bd recipes x
CV and ML exams x five models = 20 cells), under

  (a) the paper's normal approximation  mean +/- 1.96 * SE   (SE = sample sd / sqrt n,
      exactly as finetune/pairwise_analysis.py::mean_ci does), and
  (b) a percentile bootstrap over students, 2000 resamples, numpy
      default_rng(0), mirroring analysis/computer_vision_run_analysis.py::bootstrap_mae_ci
      (idx = rng.integers(0, n, size=(2000, n)); percentiles 2.5 / 97.5).

It then flags every cell whose CI upper bound lies in the open interval
(-0.05, 0.05) under either method, and states whether "7B bd, CV"
(finetune/results/X-A07bd-both-on-cv.xlsx) strictly excludes zero under each
method and what its verdict is under the Table 15 caption rule
("better" = whole 95% CI below zero; "= grader" = CI contains zero).

Data conventions (identical to finetune/pairwise_analysis.py::load_run, which is
what produced Table 15):
  CV  : AI total = "AI Total Score" (Q1-Q3 base, /35).  The workbooks'
        "TA n - Total Score (out of 35)" cells are Excel formulas whose cached
        values were dropped on rewrite (openpyxl returns None for all 114 graded
        rows), so each TA total is rebuilt as the sum of "TA n - Q1 Score",
        "TA n - Q2 Score", "TA n - Q3 Score" (Q4 has no Score column).
  ML  : AI total = "AI Total Score" + "AI Total Bonus"; TA total = sum of
        "TA n - Q1..Q3 Grade" (bonus folded in).  This equals the workbook's
        "TA n - Total Grade" column on every graded row (checked below).
  Rows: the rows with a numeric "AI Total Score"; these are exactly the pinned
        held-out students (finetune/data/eval_students.json, 114;
        finetune/data_intro/eval_students.json, 208), checked below.

Per-student statistic (Eq. paired):
  d_i = 0.5 * (|AI_i - G1_i| + |AI_i - G2_i|) - |G1_i - G2_i|

Run:  python3 analysis/checks/ftpaired_marginal_verdicts.py
"""
from __future__ import annotations

import json
import math
import statistics as st
import sys
from pathlib import Path

import numpy as np
import openpyxl
from scipy import stats

ROOT = Path(__file__).resolve().parents[2]
RESULTS = ROOT / "finetune" / "results"
EVAL_CV = ROOT / "finetune" / "data" / "eval_students.json"
EVAL_ML = ROOT / "finetune" / "data_intro" / "eval_students.json"

N_BOOT = 2000
SEED = 0
FLAG_LO, FLAG_HI = -0.05, 0.05      # open interval for "near-zero upper bound"

MODELS = [("7B", "A07"), ("14B", "A14"), ("30B", "A30M"),
          ("Gemma", "GE4B"), ("Llama", "LL8B")]
RECIPES = [("marks", ""), ("bd", "bd")]
EXAMS = [("CV", "cv"), ("ML", "intro")]


# ----------------------------------------------------------------------------
# loading (mirrors finetune/pairwise_analysis.py::load_run)
# ----------------------------------------------------------------------------
def load_rows(path: Path, exam: str):
    """-> (rows, meta): rows = list of (student_number, ai, g1, g2)."""
    ws = openpyxl.load_workbook(path, data_only=True).active
    hdr = [c.value for c in ws[1]]
    col = {h: i for i, h in enumerate(hdr) if h}
    rows = []
    meta = {"cv_total_none": 0, "cv_total_mismatch": 0, "ml_total_mismatch": 0,
            "ai_bonus_blank": 0}
    if exam == "cv":
        qs, field = (1, 2, 3), "Score"
    else:
        qs, field = (1, 2, 3), "Grade"
    for r in ws.iter_rows(min_row=2, values_only=True):
        ai = r[col["AI Total Score"]]
        if not isinstance(ai, (int, float)):
            continue
        ai = float(ai)
        if exam == "intro":
            b = r[col["AI Total Bonus"]]
            if isinstance(b, (int, float)):
                ai += float(b)
            else:
                meta["ai_bonus_blank"] += 1
        tot = {}
        for ta in (1, 2):
            v = [r[col[f"TA {ta} - Q{q} {field}"]] for q in qs]
            v = [x for x in v if isinstance(x, (int, float))]
            if len(v) == len(qs):
                tot[ta] = float(sum(v))
        if len(tot) != 2:
            continue
        if exam == "cv":
            # A formula saved without its cached result reads None here (the
            # working tree); the released tables hold the evaluated value
            # instead. Either way the totals used are rebuilt from the
            # per-question cells, and a stored total must agree with them.
            for ta in (1, 2):
                wt = r[col[f"TA {ta} - Total Score (out of 35)"]]
                if wt is None:
                    meta["cv_total_none"] += 1
                elif not isinstance(wt, (int, float)) or abs(float(wt) - tot[ta]) > 1e-9:
                    meta["cv_total_mismatch"] += 1
        else:
            for ta in (1, 2):
                wt = r[col[f"TA {ta} - Total Grade"]]
                if not isinstance(wt, (int, float)) or abs(float(wt) - tot[ta]) > 1e-9:
                    meta["ml_total_mismatch"] += 1
        rows.append((r[col["Number"]], ai, tot[1], tot[2]))
    return rows, meta


# ----------------------------------------------------------------------------
# statistics
# ----------------------------------------------------------------------------
def normal_ci(d):
    """Paper's CI: mean +/- 1.96 * sample-sd / sqrt(n)  (pairwise_analysis.mean_ci)."""
    m = st.fmean(d)
    se = st.stdev(d) / math.sqrt(len(d))
    return m, m - 1.96 * se, m + 1.96 * se, se


def t_ci(d):
    """Student-t CI (df = n-1) -- shown only as a reference, not a paper method."""
    n = len(d)
    m = st.fmean(d)
    se = st.stdev(d) / math.sqrt(n)
    tcrit = stats.t.ppf(0.975, n - 1)
    return m - tcrit * se, m + tcrit * se, tcrit


def boot_ci(d, n_boot=N_BOOT, seed=SEED, alpha=0.05):
    """Percentile bootstrap of mean(d) over students; mirrors bootstrap_mae_ci."""
    d = np.asarray(d, dtype=float)
    n = len(d)
    rng = np.random.default_rng(seed)
    idx = rng.integers(0, n, size=(n_boot, n))
    samples = d[idx].mean(axis=1)
    lo, hi = np.percentile(samples, [100 * alpha / 2, 100 * (1 - alpha / 2)])
    return float(lo), float(hi), samples


def verdict(lo, hi):
    """Caption rule: 'better' iff whole CI below 0; '= grader' iff it contains 0."""
    if lo <= 0 <= hi:
        return "= grader"
    return "worse" if lo > 0 else "better"


def fmt(x, nd=3):
    return f"{x:+.{nd}f}"


# ----------------------------------------------------------------------------
def main():
    print("# Table 15 (tab:ftpaired) recomputed to three decimals")
    print(f"# numpy {np.__version__}, scipy {__import__('scipy').__version__}, "
          f"openpyxl {openpyxl.__version__}")
    print(f"# bootstrap: {N_BOOT} resamples, numpy default_rng(seed={SEED}), "
          f"percentiles 2.5/97.5 of the resampled mean of d_i")
    print(f"# normal CI: mean +/- 1.96 * sd(ddof=1)/sqrt(n)  (finetune/pairwise_analysis.py)")
    print()

    eval_cv = set(json.load(open(EVAL_CV)))
    eval_ml = set(json.load(open(EVAL_ML)))
    print(f"held-out lists: {EVAL_CV.relative_to(ROOT)} n={len(eval_cv)}; "
          f"{EVAL_ML.relative_to(ROOT)} n={len(eval_ml)}")
    print()

    # optional cross-check against the script that produced Table 15
    pa = None
    try:
        sys.path.insert(0, str(ROOT / "finetune"))
        import pairwise_analysis as pa  # noqa: F401
        print("cross-check: imported finetune/pairwise_analysis.py (rows will be compared)")
    except Exception as e:  # pragma: no cover
        print(f"cross-check: could not import finetune/pairwise_analysis.py ({e!r}); skipped")
    print()

    cells = []
    for rec_label, rec_tag in RECIPES:
        for ex_label, ex_key in EXAMS:
            for m_label, m_tag in MODELS:
                fname = f"X-{m_tag}{rec_tag}-both-on-{ex_key}.xlsx"
                path = RESULTS / fname
                if not path.exists():
                    print(f"!! missing {path}")
                    continue
                rows, meta = load_rows(path, ex_key)
                nums = {r[0] for r in rows}
                ev = eval_cv if ex_key == "cv" else eval_ml
                assert nums == ev, (fname, len(nums), len(ev))
                if ex_key == "cv":
                    assert meta["cv_total_mismatch"] == 0, meta
                else:
                    assert meta["ml_total_mismatch"] == 0, meta
                    assert meta["ai_bonus_blank"] == 0, meta
                if pa is not None:
                    ref = pa.load_run(path)
                    mine = [(a, g1, g2) for _, a, g1, g2 in rows]
                    assert len(ref) == len(mine) and all(
                        abs(x - y) < 1e-9 for p, q in zip(ref, mine) for x, y in zip(p, q)), fname
                ai = np.array([r[1] for r in rows]); g1 = np.array([r[2] for r in rows]); g2 = np.array([r[3] for r in rows])
                ai_g1 = np.abs(ai - g1); ai_g2 = np.abs(ai - g2); g1_g2 = np.abs(g1 - g2)
                ai_gavg = np.abs(ai - (g1 + g2) / 2)
                d = 0.5 * (ai_g1 + ai_g2) - g1_g2
                dl = d.tolist()
                m, nlo, nhi, se = normal_ci(dl)
                tlo, thi, tcrit = t_ci(dl)
                blo, bhi, samples = boot_ci(dl)
                tstat, tp = stats.ttest_1samp(d, 0.0)
                cells.append(dict(
                    cell=f"{m_label} {rec_label}, {ex_label}", file=fname, n=len(rows),
                    ai_gavg=ai_gavg.mean(), ai_g1=ai_g1.mean(), ai_g2=ai_g2.mean(),
                    floor=g1_g2.mean(), dbar=m, se=se, nlo=nlo, nhi=nhi,
                    tlo=tlo, thi=thi, tcrit=tcrit, blo=blo, bhi=bhi,
                    boot_frac_ge0=float((samples >= 0).mean()), tp=tp, d=d,
                    rows=rows))

    # ------------------------------------------------------------------ table
    print("## All 20 pooled cells (three decimals)")
    print("rows used: every row of the named workbook with a numeric 'AI Total Score' "
          "(= the pinned held-out students; CV TA totals rebuilt from Q1-Q3 Score, "
          "ML from Q1-Q3 Grade)")
    print()
    hdr = (f"{'cell':18s} {'workbook':28s} {'n':>3} {'|AI-Gavg|':>9} {'|AI-G1|':>8} "
           f"{'|AI-G2|':>8} {'floor':>6} {'d-bar':>7} {'SE':>6} "
           f"{'normal 95% CI':>18} {'verdict':9} {'boot 95% CI':>18} {'verdict':9} "
           f"{'paper (2dp)':>16}")
    print(hdr); print("-" * len(hdr))
    for c in cells:
        paper = f"[{c['nlo']:+.2f}, {c['nhi']:+.2f}]"
        print(f"{c['cell']:18s} {c['file']:28s} {c['n']:>3} {c['ai_gavg']:9.3f} {c['ai_g1']:8.3f} "
              f"{c['ai_g2']:8.3f} {c['floor']:6.3f} {c['dbar']:+7.3f} {c['se']:6.3f} "
              f"[{fmt(c['nlo'])}, {fmt(c['nhi'])}] {verdict(c['nlo'], c['nhi']):9s} "
              f"[{fmt(c['blo'])}, {fmt(c['bhi'])}] {verdict(c['blo'], c['bhi']):9s} "
              f"{paper:>16}")
    print()
    print("'paper (2dp)' is the normal CI printed with the :+.2f format Table 15 uses "
          "(reproduces the table's rendering, including '-0.00').")
    print()

    # ------------------------------------------------------------- extras
    print("## Reference extras per cell (not paper methods): Student-t CI (df=n-1), "
          "paired t-test p, bootstrap share of resampled means >= 0")
    hdr2 = (f"{'cell':18s} {'n':>3} {'d-bar':>7} {'t CI (df=n-1)':>18} {'t_crit':>6} "
            f"{'t-test p':>9} {'P_boot(mean>=0)':>15}")
    print(hdr2); print("-" * len(hdr2))
    for c in cells:
        print(f"{c['cell']:18s} {c['n']:>3} {c['dbar']:+7.3f} "
              f"[{fmt(c['tlo'])}, {fmt(c['thi'])}] {c['tcrit']:6.3f} {c['tp']:9.4f} "
              f"{c['boot_frac_ge0']:15.4f}")
    print()

    # ------------------------------------------------------------- flags
    print(f"## Cells whose 95% CI upper bound lies in ({FLAG_LO}, {FLAG_HI})")
    print()
    any_flag = False
    for c in cells:
        fn = FLAG_LO < c["nhi"] < FLAG_HI
        fb = FLAG_LO < c["bhi"] < FLAG_HI
        if fn or fb:
            any_flag = True
            print(f"  {c['cell']:18s} ({c['file']})")
            print(f"      normal    upper = {c['nhi']:+.3f}  -> {'FLAGGED' if fn else 'not flagged'}; "
                  f"strictly < 0: {c['nhi'] < 0}; verdict {verdict(c['nlo'], c['nhi'])}")
            print(f"      bootstrap upper = {c['bhi']:+.3f}  -> {'FLAGGED' if fb else 'not flagged'}; "
                  f"strictly < 0: {c['bhi'] < 0}; verdict {verdict(c['blo'], c['bhi'])}")
    if not any_flag:
        print("  (none)")
    print()

    print("## Verdict changes between methods")
    changes = [c for c in cells if verdict(c['nlo'], c['nhi']) != verdict(c['blo'], c['bhi'])]
    if changes:
        for c in changes:
            print(f"  {c['cell']:18s} normal: {verdict(c['nlo'], c['nhi'])}  "
                  f"[{fmt(c['nlo'])}, {fmt(c['nhi'])}]   bootstrap: {verdict(c['blo'], c['bhi'])}  "
                  f"[{fmt(c['blo'])}, {fmt(c['bhi'])}]")
    else:
        print("  (none: every cell has the same verdict under both methods)")
    print()
    nb = sum(verdict(c['nlo'], c['nhi']) == "better" for c in cells)
    bb = sum(verdict(c['blo'], c['bhi']) == "better" for c in cells)
    print(f"count of 'better' cells: normal {nb}/20, bootstrap {bb}/20; "
          f"'worse' cells: normal {sum(verdict(c['nlo'], c['nhi'])=='worse' for c in cells)}, "
          f"bootstrap {sum(verdict(c['blo'], c['bhi'])=='worse' for c in cells)}")
    for rec in ("marks", "bd"):
        for ex in ("CV", "ML"):
            sub = [c for c in cells if c['cell'].endswith(f"{rec}, {ex}")]
            print(f"  {rec:5s} {ex}: better (normal) = "
                  f"{[c['cell'].split()[0] for c in sub if verdict(c['nlo'], c['nhi'])=='better']}; "
                  f"better (bootstrap) = "
                  f"{[c['cell'].split()[0] for c in sub if verdict(c['blo'], c['bhi'])=='better']}")
    print()

    # ------------------------------------------------------------- 7B bd CV
    c7 = next(c for c in cells if c['cell'] == "7B bd, CV")
    print("## '7B bd, CV'  (finetune/results/X-A07bd-both-on-cv.xlsx, the 114 held-out rows)")
    print(f"  n = {c7['n']};  d-bar = {c7['dbar']:+.6f};  sd(ddof=1) = {st.stdev(c7['d'].tolist()):.6f};  "
          f"SE = {c7['se']:.6f};  1.96*SE = {1.96*c7['se']:.6f}")
    print(f"  normal CI     = [{c7['nlo']:+.6f}, {c7['nhi']:+.6f}]  -> upper bound < 0: {c7['nhi'] < 0}"
          f"   (upper bound is {c7['nhi']:+.4f}; rendered '{c7['nhi']:+.2f}' at 2 dp)")
    print(f"  bootstrap CI  = [{c7['blo']:+.6f}, {c7['bhi']:+.6f}]  -> upper bound < 0: {c7['bhi'] < 0}")
    print(f"  Student-t CI  = [{c7['tlo']:+.6f}, {c7['thi']:+.6f}]  -> upper bound < 0: {c7['thi'] < 0}  (reference only)")
    print(f"  paired t-test two-sided p = {c7['tp']:.5f}; bootstrap share of resampled means >= 0: {c7['boot_frac_ge0']:.4f}")
    w = stats.wilcoxon(c7['d'], zero_method='wilcox')
    print(f"  Wilcoxon signed-rank (zeros dropped, n_eff={int((c7['d'] != 0).sum())}): "
          f"W = {w.statistic:.1f}, two-sided p = {w.pvalue:.5f}  (reference only)")
    print(f"  caption rule verdict: normal -> {verdict(c7['nlo'], c7['nhi'])};  "
          f"bootstrap -> {verdict(c7['blo'], c7['bhi'])}")
    print(f"  per-student d_i: min {c7['d'].min():+.3f}, median {np.median(c7['d']):+.3f}, "
          f"max {c7['d'].max():+.3f}; #d<0 {(c7['d']<0).sum()}, #d=0 {(c7['d']==0).sum()}, #d>0 {(c7['d']>0).sum()}")
    print()

    # bootstrap seed sensitivity for every flagged / marginal cell (|upper| < 0.10 under either method)
    print("## Bootstrap seed sensitivity (2000 resamples each) for cells with |upper bound| < 0.10 under either method")
    print("   (seed 0 is the reported one; the others show how far the percentile upper bound moves with the RNG stream)")
    for c in cells:
        if min(abs(c['nhi']), abs(c['bhi'])) < 0.10:
            ups = []
            for s in range(10):
                lo, hi, _ = boot_ci(c['d'].tolist(), seed=s)
                ups.append(hi)
            lo20k, hi20k, _ = boot_ci(c['d'].tolist(), n_boot=20000, seed=0)
            print(f"  {c['cell']:18s} normal upper {c['nhi']:+.3f} | boot upper by seed 0..9: "
                  + ", ".join(f"{u:+.4f}" for u in ups)
                  + f" | seeds with upper >= 0: {sum(u >= 0 for u in ups)}/10"
                  + f" | 20000-resample seed 0: [{lo20k:+.3f}, {hi20k:+.3f}]")
    print()

    # ------------------------------------------------------------- floors
    print("## Human floors on the held-out students (mean |G1-G2|)")
    for ex in ("CV", "ML"):
        c = next(c for c in cells if c['cell'].endswith(f"marks, {ex}"))
        print(f"  {ex}: {c['floor']:.3f}  (paper caption: {'2.83' if ex=='CV' else '5.19'})")
    print()

    # ------------------------------------------------------------- LaTeX rows
    print("## Table 15 rows re-rendered at three decimals (normal CI, the paper's method)")
    for c in cells:
        v = "\\textbf{better}" if verdict(c['nlo'], c['nhi']) == "better" else "$=$ grader"
        print(f"{c['cell']:16s} & {c['n']} & {c['ai_gavg']:.2f} & {c['ai_g1']:.2f} & "
              f"${c['dbar']:+.2f}$ $[{c['nlo']:+.3f}, {c['nhi']:+.3f}]$ & {v} \\\\")
    print()
    print("## Same rows with the bootstrap CI (2000 resamples, seed 0)")
    for c in cells:
        v = "\\textbf{better}" if verdict(c['blo'], c['bhi']) == "better" else "$=$ grader"
        print(f"{c['cell']:16s} & {c['n']} & {c['ai_gavg']:.2f} & {c['ai_g1']:.2f} & "
              f"${c['dbar']:+.2f}$ $[{c['blo']:+.3f}, {c['bhi']:+.3f}]$ & {v} \\\\")


if __name__ == "__main__":
    main()
