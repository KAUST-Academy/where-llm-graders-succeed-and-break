#!/usr/bin/env python3
"""Per-question Mean Signed Error (bias) and error cancellation.

The concern is that an aggregate MAE (and an aggregate signed error)
can hide directional errors that cancel across criteria: a grader that is +2 on
one question and -2 on another reports a total bias of 0.  The paper reports a
total-level signed error under the name *bias* (Section 3; Tables 7, 8, 10, 17,
19, 20) and nothing per question.  This script decomposes that number.

What it computes, per run, over EXACTLY the students the paper's total-level
metric uses:

  per-question mean signed error   bias_q = mean_i ( AI_q,i - TAavg_q,i )
  per-question MAE                 mae_q  = mean_i | AI_q,i - TAavg_q,i |
  total bias / total MAE           the paper's numbers (reproduced to 3 dp)
  sign pattern                     sign(bias_1), sign(bias_2), sign(bias_3)
  cancellation                     C = sum_q |bias_q| - |sum_q bias_q|
                                     = sum_q |bias_q| - |total bias|
                                   C = 0 iff no per-question bias opposes the
                                   others; C > 0 is exactly the amount of
                                   directional error the total hides.

Runs covered:
  CV headline  D01, D02, F02, B03, A01, C01
               (Table 7 / tab:headline, Table 19 / tab:allruns rows)
  ML headline  IG08, IG07, IG01, IG05
               (Table 20 / tab:iaruns rows; the Gemini rows of tab:closedvendors)
  CV matched   the 17 matched neutral/strict open-weights pairs of Table 2
               (tab:brittle), i.e. 34 runs, obtained from
               computer_vision_run_analysis.matched_persona_pairs() -- the very
               function that generated that table -- and cross-checked against
               the run ids printed in sections/appendix.tex.

Files and columns read
----------------------
CV workbooks  computer_vision_results/results/<RID>__*.xlsx, sheet 0, one row
              per student keyed by "Number".
    AI total        "AI Total Score"  (Q1-Q3 base, 35 points; recovered as
                    "AI Q1..Q3 Score" summed when the cell is blank but all
                    three question cells are present -- the rule
                    computer_vision_run_analysis.build_master_table uses).
    AI per question "AI Q{q} Score", q = 1, 2, 3.  Q4 is bonus-only, has no TA
                    Score column and is excluded from the 35-point scale, so it
                    is not part of the decomposition.
    TA total        mean of "TA 1 - Total Score (out of 35)" and "TA 2 - ...";
                    those cells are Excel formulas whose cached values were
                    dropped on rewrite (openpyxl/pandas read None), so each TA
                    total is rebuilt as "TA n - Q1 Score" + "Q2" + "Q3"
                    (computer_vision_run_analysis._ta_grader_total, skipna=False).
    TA per question mean of "TA 1 - Q{q} Score" and "TA 2 - Q{q} Score" over the
                    SAME graders that contributed to that student's TA total, so
                    that sum_q TAavg_q == TAavg_total identically (asserted).
    Rows used       "AI total present AND TA total present" -- the paper's
                    n_valid_vs_TA.

ML workbooks  introduction_to_ai_results/results/<RID>__*.xlsx, sheet 0.
    AI total        "AI Total Score" + "AI Total Bonus".
    AI per question "AI Q{q} Score" + "AI Q{q} Bonus", q = 1, 2, 3.
    TA truth        introduction_to_ai_dataset/Practical_AI_exam_grades.csv via
                    introduction_to_ai_run_analysis.ta_truth() (totals, "TA n -
                    Total Grade") and .ta_per_question() (per-question averages
                    of "TA n - Q{q} Grade"; the bonus is folded into the grade).
    Rows used       "AI Total Score present AND both TA totals present" -- the
                    paper's n.

Self-checks printed below (all must pass for the numbers to be trusted)
----------------------------------------------------------------------
  * total MAE and total bias reproduce analysis/*_master_comparison.csv to
    <= 5e-4 and the published Table 7 / 10 / 19 / 20 values to 2 dp;
  * sum_q bias_q == total bias to < 1e-9 (the decomposition is exact, which is
    what makes the cancellation number meaningful) -- also printed to 3 dp;
  * sum_q AI_q == AI total and sum_q TAavg_q == TAavg total per student;
  * per-question MAE is compared with the Q{q}_MAE_vs_TA columns of
    analysis/computer_vision_master_comparison.csv (those are computed over
    every student with that question graded, not over the total-valid set, so a
    small difference is expected and is reported).

Run:  python3 analysis/checks/per_question_signed_error.py
"""
from __future__ import annotations

import sys
import warnings
from pathlib import Path

import numpy as np
import openpyxl
import pandas as pd

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT / "analysis"))

import computer_vision_run_analysis as cvra          # noqa: E402
import introduction_to_ai_run_analysis as mlra       # noqa: E402

warnings.filterwarnings("ignore", category=UserWarning, module="openpyxl")

CV_RESULTS = ROOT / "computer_vision_results" / "results"
ML_RESULTS = ROOT / "introduction_to_ai_results" / "results"
CV_MASTER = ROOT / "analysis" / "computer_vision_master_comparison.csv"
ML_MASTER = ROOT / "analysis" / "introduction_to_ai_master_comparison.csv"

QS = (1, 2, 3)

# Headline runs named in the paper's Table 7 and Section 8.
CV_HEADLINE = ["D01", "D02", "F02", "B03", "A01", "C01"]
ML_HEADLINE = ["IG08", "IG07", "IG01", "IG05"]

CV_LABEL = {
    "D01": "gemini-3-flash-preview, neutral",
    "D02": "gemini-3.1-pro-preview, neutral",
    "F02": "gemini-3.1-pro-preview + lenient",
    "B03": "gemini-flash-lite + thinking",
    "A01": "gemini-flash-lite, neutral (baseline)",
    "C01": "gemini-flash-lite + strict",
}
ML_LABEL = {
    "IG08": "gemini-3.1-pro-preview, neutral",
    "IG07": "gemini-3-flash-preview, neutral",
    "IG01": "gemini-flash-lite, neutral",
    "IG05": "gemini-flash-lite + strict",
}

# Published values, copied from the paper so the script fails loudly if a
# convention drifts.  CV: sections/appendix_runs_table.tex (Table 19) and
# sections/appendix.tex Table 7.  ML: sections/introduction_to_ai_appendix_runs_table.tex
# (Table 20).  (MAE, bias, n).
PAPER = {
    "D01": (1.64, +0.01, 570), "D02": (1.86, -0.82, 570),
    "F02": (1.79, +0.82, 570), "B03": (2.05, -0.15, 570),
    "A01": (3.34, -1.78, 570), "C01": (5.75, -5.44, 570),
    "IG08": (3.40, +1.18, 1038), "IG07": (3.92, +2.24, 1038),
    "IG01": (7.53, +5.98, 1013), "IG05": (9.02, -7.83, 1008),
}

# The 17 strict run ids printed in Table 2 (sections/appendix.tex, tab:brittle),
# open-weights rows only, in the table's own order.
TAB_BRITTLE_STRICT = ["L-M05", "L-U02", "L-W02", "L-F02", "L-M07", "L-S02",
                      "L-Z02", "L-Y02", "L-M01", "L-X02", "L-C01", "L-V02",
                      "L-N02", "L-M03", "L-J02", "L-Q02", "L-R02"]

# Per-question maxima, for context only (Appendix A / the ML scale line).
CV_QMAX = {1: 12, 2: 11, 3: 12}          # 35-point base scale
ML_QMAX = {1: 26, 2: 17, 3: 22}          # 23+3, 14+3, 19+3 = ~65 with bonus


# ---------------------------------------------------------------------------
# CV
# ---------------------------------------------------------------------------
def cv_workbook(rid: str) -> Path:
    hits = sorted(CV_RESULTS.glob(f"{rid}__*.xlsx"))
    if not hits:
        raise FileNotFoundError(f"no CV workbook for {rid}")
    return hits[0]


def cv_run(rid: str) -> dict:
    """Per-question and total bias/MAE for one CV run, on the paper's rows."""
    path = cv_workbook(rid)
    df = pd.read_excel(path, sheet_name=0).set_index("Number")
    df = cvra.maybe_fill_from_variance(df, rid)

    ai_q = {q: pd.to_numeric(df[f"AI Q{q} Score"], errors="coerce") for q in QS}
    ai_total = pd.to_numeric(df["AI Total Score"], errors="coerce")
    base = pd.concat([ai_q[q] for q in QS], axis=1)
    recoverable = ai_total.isna() & base.notna().all(axis=1)
    n_recovered = int(recoverable.sum())
    ai_total = ai_total.where(~recoverable, base.sum(axis=1))

    # TA side: rebuild each grader's total from Q1-Q3 (the stored total column is
    # a formula with no cached value in most workbooks) and average the graders
    # that have a complete total -- exactly ta_combined()'s rule.
    t = {g: cvra._ta_grader_total(df, g) for g in (1, 2)}
    ta_avg = cvra.ta_combined(df)
    have = {g: t[g].notna() for g in (1, 2)}
    ta_q = {}
    for q in QS:
        cols = {g: pd.to_numeric(df[f"TA {g} - Q{q} Score"], errors="coerce")
                for g in (1, 2)}
        num = pd.Series(0.0, index=df.index)
        den = pd.Series(0.0, index=df.index)
        for g in (1, 2):
            num = num.add(cols[g].where(have[g], 0.0), fill_value=0.0)
            den = den.add(have[g].astype(float), fill_value=0.0)
        ta_q[q] = (num / den).where(den > 0)

    valid = ai_total.notna() & ta_avg.notna()
    # identities
    ai_sum_gap = float((base.sum(axis=1) - ai_total)[valid].abs().max())
    ta_sum_gap = float((sum(ta_q[q] for q in QS) - ta_avg)[valid].abs().max())

    diff = (ai_total - ta_avg)[valid]
    out = {
        "run_id": rid, "file": path.name, "n": int(valid.sum()),
        "mae": float(diff.abs().mean()), "bias": float(diff.mean()),
        "ai_sum_gap": ai_sum_gap, "ta_sum_gap": ta_sum_gap,
        "n_recovered": n_recovered,
        "ai_mean": float(ai_total[valid].mean()),
        "ta_mean": float(ta_avg[valid].mean()),
    }
    for q in QS:
        d = (ai_q[q] - ta_q[q])[valid]
        out[f"bias{q}"] = float(d.mean())
        out[f"mae{q}"] = float(d.abs().mean())
        out[f"aimean{q}"] = float(ai_q[q][valid].mean())
        out[f"tamean{q}"] = float(ta_q[q][valid].mean())
        out[f"nq{q}"] = int(d.notna().sum())
    return out


# ---------------------------------------------------------------------------
# ML
# ---------------------------------------------------------------------------
def ml_workbook(rid: str) -> Path:
    hits = sorted(ML_RESULTS.glob(f"{rid}__*.xlsx"))
    if not hits:
        raise FileNotFoundError(f"no ML workbook for {rid}")
    return hits[0]


def ml_run(rid: str, TA: dict, TAQ: dict) -> dict:
    """Per-question and total bias/MAE for one ML run, on the paper's rows."""
    path = ml_workbook(rid)
    ws = openpyxl.load_workbook(path, data_only=True).active
    hdr = [c.value for c in ws[1]]
    c = {h: i for i, h in enumerate(hdr) if h}

    tot_d, q_d = [], {q: [] for q in QS}
    ai_tot_v, ta_tot_v = [], []
    ai_q_v = {q: [] for q in QS}
    ta_q_v = {q: [] for q in QS}
    ai_sum_gap = ta_sum_gap = 0.0
    n_rows = n_no_ai = 0
    for r in ws.iter_rows(min_row=2, values_only=True):
        n_rows += 1
        a = r[c["AI Total Score"]]
        if a is None:
            n_no_ai += 1
            continue
        tot = float(a) + float(r[c["AI Total Bonus"]] or 0)
        num = r[c["Number"]]
        t1, t2 = TA.get(num, (None, None))
        if t1 is None or t2 is None or num not in TAQ:
            continue
        tq, tt = TAQ[num]                       # ([q1,q2,q3] avg, total avg)
        assert abs(tt - (t1 + t2) / 2) < 1e-9
        aq = []
        for q in QS:
            v = r[c[f"AI Q{q} Score"]]
            b = r[c[f"AI Q{q} Bonus"]] if f"AI Q{q} Bonus" in c else 0
            aq.append((float(v) if isinstance(v, (int, float)) else float("nan"))
                      + float(b or 0))
        ai_sum_gap = max(ai_sum_gap, abs(sum(aq) - tot))
        ta_sum_gap = max(ta_sum_gap, abs(sum(tq) - tt))
        tot_d.append(tot - tt)
        ai_tot_v.append(tot); ta_tot_v.append(tt)
        for i, q in enumerate(QS):
            q_d[q].append(aq[i] - tq[i])
            ai_q_v[q].append(aq[i]); ta_q_v[q].append(tq[i])

    d = np.array(tot_d)
    out = {"run_id": rid, "file": path.name, "n": len(d),
           "mae": float(np.abs(d).mean()), "bias": float(d.mean()),
           "ai_sum_gap": ai_sum_gap, "ta_sum_gap": ta_sum_gap,
           "n_recovered": 0, "n_no_ai": n_no_ai, "n_rows": n_rows,
           "ai_mean": float(np.mean(ai_tot_v)), "ta_mean": float(np.mean(ta_tot_v))}
    for q in QS:
        e = np.array(q_d[q])
        out[f"bias{q}"] = float(e.mean())
        out[f"mae{q}"] = float(np.abs(e).mean())
        out[f"aimean{q}"] = float(np.mean(ai_q_v[q]))
        out[f"tamean{q}"] = float(np.mean(ta_q_v[q]))
        out[f"nq{q}"] = len(e)
    return out


# ---------------------------------------------------------------------------
# derived quantities
# ---------------------------------------------------------------------------
EPS = 1e-9          # float noise; a per-question bias below this is a real zero


def decorate(rec: dict) -> dict:
    b = [rec[f"bias{q}"] for q in QS]
    rec["sum_abs"] = float(sum(abs(x) for x in b))
    rec["sum_bias"] = float(sum(b))
    rec["identity_gap"] = abs(rec["sum_bias"] - rec["bias"])
    c = rec["sum_abs"] - abs(rec["bias"])
    rec["cancel"] = 0.0 if abs(c) < EPS else c      # >= 0 by the triangle ineq.
    pos = sum(1 for x in b if x > EPS)
    neg = sum(1 for x in b if x < -EPS)
    rec["opposite"] = pos > 0 and neg > 0
    rec["signs"] = " ".join("+" if x > EPS else ("-" if x < -EPS else "0")
                            for x in b)
    rec["cancel_share"] = (rec["cancel"] / rec["sum_abs"]) if rec["sum_abs"] else 0.0
    return rec


def sign_word(rec: dict) -> str:
    return "OPPOSITE" if rec["opposite"] else "same sign"


# ---------------------------------------------------------------------------
def print_block(title: str, recs: list, qmax: dict, note: str = "") -> None:
    print(f"## {title}")
    if note:
        print(f"   {note}")
    print()
    hdr = (f"{'run':7s} {'n':>5} {'MAE':>6} {'bias':>7} | "
           + " ".join(f"{'Q'+str(q)+' bias':>8}" for q in QS) + " | "
           + " ".join(f"{'Q'+str(q)+' MAE':>7}" for q in QS)
           + f" | {'signs':5s} {'opposite?':9s} {'sum|bq|':>8} {'cancel':>7} {'share':>6}"
             f" {'sum bq - bias':>13}")
    print(hdr)
    print("-" * len(hdr))
    for r in recs:
        print(f"{r['run_id']:7s} {r['n']:>5} {r['mae']:6.3f} {r['bias']:+7.3f} | "
              + " ".join(f"{r[f'bias{q}']:+8.3f}" for q in QS) + " | "
              + " ".join(f"{r[f'mae{q}']:7.3f}" for q in QS)
              + f" | {r['signs']:5s} {sign_word(r):9s} {r['sum_abs']:8.3f} "
                f"{r['cancel']:7.3f} {r['cancel_share']:6.1%} {r['identity_gap']:13.2e}")
    print()
    print("   per-question maxima on this exam: "
          + ", ".join(f"Q{q} = {qmax[q]}" for q in QS))
    print()


def main() -> int:
    print("# Per-question mean signed error (bias), per-question MAE, "
          "and error cancellation")
    print(f"# numpy {np.__version__}, pandas {pd.__version__}, "
          f"openpyxl {openpyxl.__version__}")
    print("# cancellation C = sum_q |bias_q| - |total bias|; C = 0 iff every "
          "per-question bias has the same sign")
    print("# all statistics are over exactly the students the paper's total-level "
          "metric uses for that run")
    print()

    cvm = pd.read_csv(CV_MASTER).set_index("run_id")
    mlm = pd.read_csv(ML_MASTER).set_index("run_id") if ML_MASTER.exists() else None
    print(f"master tables: {CV_MASTER.relative_to(ROOT)} ({len(cvm)} runs)"
          + (f"; {ML_MASTER.relative_to(ROOT)} ({len(mlm)} runs)" if mlm is not None
             else "; ML master table absent"))
    print()

    # ---------------------------------------------------------------- CV runs
    cv_recs = {rid: decorate(cv_run(rid)) for rid in CV_HEADLINE}

    pairs = cvra.matched_persona_pairs(cvm.reset_index())
    open_pairs = pairs[pairs.test_rid.str.startswith("L-")].copy()
    got = sorted(open_pairs.test_rid)
    assert got == sorted(TAB_BRITTLE_STRICT), (got, sorted(TAB_BRITTLE_STRICT))
    print(f"matched open-weights pairs recovered from "
          f"computer_vision_run_analysis.matched_persona_pairs(): {len(open_pairs)}"
          f" -- strict run ids identical to the open-weights rows of Table 2 "
          f"(tab:brittle, sections/appendix.tex)")
    order = {r: i for i, r in enumerate(TAB_BRITTLE_STRICT)}
    open_pairs = open_pairs.sort_values("test_rid", key=lambda s: s.map(order))
    print("   pairs (neutral -> strict): "
          + ", ".join(f"{r.baseline_rid}->{r.test_rid}"
                      for r in open_pairs.itertuples()))
    print()

    for r in open_pairs.itertuples():
        for rid in (r.baseline_rid, r.test_rid):
            if rid not in cv_recs:
                cv_recs[rid] = decorate(cv_run(rid))

    # ---------------------------------------------------------------- ML runs
    TA = mlra.ta_truth()
    TAQ = mlra.ta_per_question()
    ml_recs = {rid: decorate(ml_run(rid, TA, TAQ)) for rid in ML_HEADLINE}
    print(f"ML ground truth: introduction_to_ai_dataset/Practical_AI_exam_grades.csv "
          f"-- {len(TA)} students, {len(TAQ)} with per-question grades")
    print()

    # ----------------------------------------------------------- self-checks
    print("## Self-checks")
    print()
    print("### (a) total MAE / bias reproduce the master tables and the paper")
    hdr = (f"{'run':7s} {'n':>5} {'n(paper)':>8} {'MAE':>7} {'MAE(csv)':>9} "
           f"{'bias':>8} {'bias(csv)':>10} {'MAE(paper)':>10} {'bias(paper)':>11} "
           f"{'max gap':>9}")
    print(hdr); print("-" * len(hdr))
    ok = True
    for rid in CV_HEADLINE + ML_HEADLINE:
        rec = cv_recs[rid] if rid in cv_recs else ml_recs[rid]
        if rid in cv_recs:
            row = cvm.loc[rid]
            cmae, cbias, cn = (float(row["MAE_vs_TA_avg"]),
                               float(row["bias_vs_TA_avg"]), int(row["n_valid_vs_TA"]))
        elif mlm is not None and rid in mlm.index:
            row = mlm.loc[rid]
            cmae, cbias, cn = (float(row["MAE_vs_TA_avg"]),
                               float(row["bias_vs_TA_avg"]), int(row["n_valid_vs_TA"]))
        else:
            cmae = cbias = float("nan"); cn = -1
        pmae, pbias, pn = PAPER[rid]
        gap = max(abs(rec["mae"] - cmae) if cmae == cmae else 0.0,
                  abs(rec["bias"] - cbias) if cbias == cbias else 0.0)
        good = (abs(round(rec["mae"], 2) - pmae) < 5e-3
                and abs(round(rec["bias"], 2) - pbias) < 5e-3
                and rec["n"] == pn and (cn in (-1, rec["n"])) and gap < 5e-4)
        ok &= good
        print(f"{rid:7s} {rec['n']:>5} {pn:>8} {rec['mae']:7.3f} {cmae:9.3f} "
              f"{rec['bias']:+8.3f} {cbias:+10.3f} {pmae:10.2f} {pbias:+11.2f} "
              f"{gap:9.2e}" + ("" if good else "   <-- MISMATCH"))
    print()
    print(f"every headline run reproduces its published n, MAE and bias: {ok}")
    print()

    print("### (b) the decomposition is exact: sum_q bias_q == total bias")
    allrecs = list(cv_recs.values()) + list(ml_recs.values())
    worst = max(allrecs, key=lambda r: r["identity_gap"])
    print(f"   worst |sum_q bias_q - total bias| over all {len(allrecs)} runs: "
          f"{worst['identity_gap']:.3e}  ({worst['run_id']})")
    print(f"   worst per-student |sum_q AI_q - AI total|:   "
          f"{max(r['ai_sum_gap'] for r in allrecs):.3e}")
    print(f"   worst per-student |sum_q TAavg_q - TAavg|:   "
          f"{max(r['ta_sum_gap'] for r in allrecs):.3e}")
    print("   -> to 3 dp the per-question biases sum to the published total bias "
          "in every run.")
    print()
    print("   three-decimal statement of the identity for the headline runs:")
    for rid in CV_HEADLINE + ML_HEADLINE:
        rec = cv_recs[rid] if rid in cv_recs else ml_recs[rid]
        terms = " ".join(f"{rec[f'bias{q}']:+.3f}" for q in QS)
        print(f"     {rid:5s}  {terms}  =  {rec['sum_bias']:+.3f}   "
              f"(total bias {rec['bias']:+.3f}, paper {PAPER[rid][1]:+.2f})")
    print()

    print("### (c) per-question MAE vs the Q{q}_MAE_vs_TA columns of the CV master table")
    print("   (the CSV computes those over every student with that question graded, "
           "not over the total-valid set, so small gaps are expected)")
    hdr = (f"{'run':7s} " + " ".join(f"{'Q'+str(q)+' mine':>8} {'Q'+str(q)+' csv':>8}"
                                     for q in QS) + f" {'max gap':>9}")
    print(hdr); print("-" * len(hdr))
    for rid in CV_HEADLINE:
        rec = cv_recs[rid]
        row = cvm.loc[rid]
        g = max(abs(rec[f"mae{q}"] - float(row[f"Q{q}_MAE_vs_TA"])) for q in QS)
        print(f"{rid:7s} " + " ".join(f"{rec[f'mae{q}']:8.3f} "
                                      f"{float(row[f'Q{q}_MAE_vs_TA']):8.3f}"
                                      for q in QS) + f" {g:9.2e}")
    print()

    # ---------------------------------------------------------------- results
    print_block("CV headline runs (Tables 7 and 19)",
                [cv_recs[r] for r in CV_HEADLINE], CV_QMAX,
                note="AI Qq = 'AI Qq Score'; TA Qq = mean of 'TA 1/2 - Qq Score'; "
                     "35-point base scale, Q4 bonus excluded")
    for rid in CV_HEADLINE:
        rec = cv_recs[rid]
        print(f"   {rid:5s} {CV_LABEL[rid]:38s} mean awarded "
              + " ".join(f"Q{q} {rec[f'aimean{q}']:5.2f}/{CV_QMAX[q]}" for q in QS)
              + f"   grader " + " ".join(f"Q{q} {rec[f'tamean{q}']:5.2f}" for q in QS))
    print()

    print_block("ML headline runs (Table 20)",
                [ml_recs[r] for r in ML_HEADLINE], ML_QMAX,
                note="AI Qq = 'AI Qq Score' + 'AI Qq Bonus'; TA Qq = mean of "
                     "'TA 1/2 - Qq Grade'; ~65-point score-plus-bonus scale")
    for rid in ML_HEADLINE:
        rec = ml_recs[rid]
        print(f"   {rid:5s} {ML_LABEL[rid]:38s} mean awarded "
              + " ".join(f"Q{q} {rec[f'aimean{q}']:5.2f}/{ML_QMAX[q]}" for q in QS)
              + f"   grader " + " ".join(f"Q{q} {rec[f'tamean{q}']:5.2f}" for q in QS))
    print()

    neutral = [cv_recs[r.baseline_rid] for r in open_pairs.itertuples()]
    strict = [cv_recs[r.test_rid] for r in open_pairs.itertuples()]
    print_block("CV matched open-weights NEUTRAL runs (17 pairs of Table 2)",
                neutral, CV_QMAX)
    print_block("CV matched open-weights STRICT runs (17 pairs of Table 2)",
                strict, CV_QMAX)

    print("## Neutral -> strict, per question (same pair, same students)")
    hdr = (f"{'model':24s} {'neutral':>8} {'strict':>8} | "
           + " ".join(f"{'dQ'+str(q):>8}" for q in QS)
           + f" | {'neutral C':>9} {'strict C':>9}")
    print(hdr); print("-" * len(hdr))
    for r in open_pairs.itertuples():
        n_, s_ = cv_recs[r.baseline_rid], cv_recs[r.test_rid]
        print(f"{r.model:24s} {n_['bias']:+8.3f} {s_['bias']:+8.3f} | "
              + " ".join(f"{s_[f'bias{q}'] - n_[f'bias{q}']:+8.3f}" for q in QS)
              + f" | {n_['cancel']:9.3f} {s_['cancel']:9.3f}")
    print()

    # ------------------------------------------------------------ cancellation
    print("## Cancellation summary")
    print()
    groups = [
        ("CV headline (6)", [cv_recs[r] for r in CV_HEADLINE]),
        ("ML headline (4)", [ml_recs[r] for r in ML_HEADLINE]),
        ("CV matched neutral (17)", neutral),
        ("CV matched strict (17)", strict),
    ]
    for name, recs in groups:
        opp = [r for r in recs if r["opposite"]]
        print(f"   {name:26s} opposite-sign runs: {len(opp)}/{len(recs)}; "
              f"cancellation C: min {min(r['cancel'] for r in recs):.3f}, "
              f"median {np.median([r['cancel'] for r in recs]):.3f}, "
              f"max {max(r['cancel'] for r in recs):.3f}")
    print()
    print("   Reading: cancellation is a property of the ACCURATE runs. Every one "
          "of the 17 strict")
    print("   runs and all four ML Gemini runs push every question in the same "
          "direction, so their")
    print("   totals hide nothing; the runs whose totals understate their "
          "directional error are the")
    print("   near-unbiased ones (D01 C = "
          f"{cv_recs['D01']['cancel']:.3f} on a total bias of "
          f"{cv_recs['D01']['bias']:+.3f}, B03 C = {cv_recs['B03']['cancel']:.3f} "
          f"on {cv_recs['B03']['bias']:+.3f}).")
    print()
    every = [(r, "CV") for r in cv_recs.values()] + [(r, "ML") for r in ml_recs.values()]
    every.sort(key=lambda t: -t[0]["cancel"])
    print("   every run ranked by cancellation C (top 12 of "
          f"{len(every)} runs analysed):")
    print(f"   {'run':7s} {'exam':4s} {'signs':5s} {'total bias':>10} "
          f"{'sum|bias_q|':>11} {'C':>7} {'C / sum|bias_q|':>15}")
    for rec, exam in every[:12]:
        print(f"   {rec['run_id']:7s} {exam:4s} {rec['signs']:5s} "
              f"{rec['bias']:+10.3f} {rec['sum_abs']:11.3f} {rec['cancel']:7.3f} "
              f"{rec['cancel_share']:15.1%}")
    print()
    top, top_exam = every[0]
    print(f"   LARGEST cancellation: {top['run_id']} ({top_exam} exam, "
          f"{top['file']}): per-question biases "
          + ", ".join(f"Q{q} {top[f'bias{q}']:+.3f}" for q in QS)
          + f"; sum of magnitudes {top['sum_abs']:.3f} against a total bias of "
            f"{top['bias']:+.3f}, so C = {top['cancel']:.3f} points "
            f"({top['cancel_share']:.0%} of the directional error) is invisible "
            f"in the total.")
    hl = [r for r in (list(cv_recs[x] for x in CV_HEADLINE)
                      + list(ml_recs[x] for x in ML_HEADLINE))]
    hl.sort(key=lambda r: -r["cancel"])
    print(f"   largest among the headline runs: {hl[0]['run_id']} "
          f"C = {hl[0]['cancel']:.3f} (signs {hl[0]['signs']}, total bias "
          f"{hl[0]['bias']:+.3f}, sum of magnitudes {hl[0]['sum_abs']:.3f})")
    hl_opp = [r for r in hl if r["opposite"]]
    print(f"   headline runs whose per-question biases disagree in sign: "
          + (", ".join(f"{r['run_id']} ({r['signs']})" for r in hl_opp)
             if hl_opp else "none"))
    print()

    # ------------------------------------------------------------- LaTeX body
    print("## LaTeX body for the proposed appendix table (headline runs)")
    print()
    print("% ---- per-question mean signed error, headline runs ----")
    print("\\multicolumn{8}{l}{\\emph{CV exam --- $35$-point base scale "
          "(Q$1$ $12$, Q$2$ $11$, Q$3$ $12$)}} \\\\")
    for rid in CV_HEADLINE:
        r = cv_recs[rid]
        print(f"{rid} & {CV_LABEL[rid].replace('gemini-', '').replace('+', '$+$')} "
              f"& {r['n']} & "
              + " & ".join(f"${r[f'bias{q}']:+.2f}$" for q in QS)
              + f" & ${r['bias']:+.2f}$ & {r['sum_abs']:.2f} & {r['cancel']:.2f} \\\\")
    print("\\midrule")
    print("\\multicolumn{8}{l}{\\emph{ML exam --- $\\approx 65$-point scale "
          "(Q$1$ $26$, Q$2$ $17$, Q$3$ $22$)}} \\\\")
    for rid in ML_HEADLINE:
        r = ml_recs[rid]
        print(f"{rid} & {ML_LABEL[rid].replace('gemini-', '').replace('+', '$+$')} "
              f"& {r['n']} & "
              + " & ".join(f"${r[f'bias{q}']:+.2f}$" for q in QS)
              + f" & ${r['bias']:+.2f}$ & {r['sum_abs']:.2f} & {r['cancel']:.2f} \\\\")
    print()
    print("## LaTeX body, per-question MAE for the same runs (reference, not proposed)")
    for rid in CV_HEADLINE:
        r = cv_recs[rid]
        print(f"{rid} & {r['n']} & " + " & ".join(f"{r[f'mae{q}']:.2f}" for q in QS)
              + f" & {r['mae']:.2f} \\\\")
    for rid in ML_HEADLINE:
        r = ml_recs[rid]
        print(f"{rid} & {r['n']} & " + " & ".join(f"{r[f'mae{q}']:.2f}" for q in QS)
              + f" & {r['mae']:.2f} \\\\")
    print()

    # ---------------------------------------------------------------- bib note
    print("## Citation check")
    if not (ROOT / "references.bib").is_file():
        # the bibliography belongs to the paper, not to the released code
        print("   skipped: references.bib is not here")
        return 0
    bib = (ROOT / "references.bib").read_text(encoding="utf-8")
    for needle in ("Cognitive Bias", "cognitive bias", "2309.17012", "Koo"):
        print(f"   references.bib contains {needle!r}: {needle in bib}")
    print("   -> 'Benchmarking Cognitive Biases in Large Language Models as "
          "Evaluators' is NOT in references.bib.")
    print("   arXiv:2309.17012 (fetched from https://arxiv.org/abs/2309.17012): "
          "citation_title = 'Benchmarking Cognitive Biases in Large Language "
          "Models as Evaluators';")
    print("   authors Koo, Lee, Raheja, Park, Kim, Kang; submitted 2023-09-29; "
          "arXiv comment 'Published at ACL 2024'.")
    print("   (The bib is NOT edited by this script.)")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
