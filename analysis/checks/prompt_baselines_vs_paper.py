#!/usr/bin/env python3
"""Prompt-only baselines for the strict-persona collapse.

Compares the two non-parametric prompt strategies run by prompt_baselines_grade.py
(``decomposed``: one call per rubric row; ``arbitrated``: strict grade + neutral
grade + neutral arbiter) against the paper's own runs on the SAME students
(CV exam, students 1-100 unless a run graded fewer), for the four collapsing
open-weights models named in Section 5:

    PB-?01  Qwen2.5-Coder-32B    paper strict L-C01 (20.29, selective)   neutral L-A32
    PB-?02  GLM-4-32B            paper strict L-X02 ( 9.66, selective)   neutral L-X01
    PB-?03  Mistral-Small-24B    paper strict L-Z02 (26.03, refusal)     neutral L-Z01
    PB-?04  Llama-3.1-8B         paper strict L-U02 (26.04, refusal)     neutral L-U01

Metric conventions are the paper's (Section 3): AI total = AI Q1+Q2+Q3 base
scores (35 points) vs the mean of the two graders' Q1-Q3 totals rebuilt from
the per-question "TA n - Qk Score" columns (the workbooks' total formulas
carry no cached values after a rewrite); rows with any blank cell excluded.
MAE = mean |AI - TA_avg|, bias = mean (AI - TA_avg), 95% CI = percentile
bootstrap over students, 2000 resamples, numpy default_rng(0). Behaviour
classes as in Section 5.1 (refusal >= 90% zero totals and sd < 0.5; near-
refusal >= 90% zeros; collapse MAE >= 8 on this exam; graded otherwise).
The paper runs are restricted to exactly the students the baseline graded, so
every comparison is paired on students; the human floor on the same students
is reported alongside.

Nothing is written to the paper; the script prints every number.

    python3 analysis/checks/prompt_baselines_vs_paper.py            # all four models, both modes
    python3 analysis/checks/prompt_baselines_vs_paper.py --models 01 02
"""
from __future__ import annotations

import argparse
import glob
import statistics
import sys
from pathlib import Path

import numpy as np
import openpyxl

ROOT = Path(__file__).resolve().parents[2]
PB_RESULTS = ROOT / "prompt_baselines" / "results"
PAPER_RESULTS = ROOT / "computer_vision_results" / "results"

MODELS = {
    "01": ("Qwen2.5-Coder-32B", "L-C01", "L-A32"),
    "02": ("GLM-4-32B", "L-X02", "L-X01"),
    "03": ("Mistral-Small-24B", "L-Z02", "L-Z01"),
    "04": ("Llama-3.1-8B", "L-U02", "L-U01"),
}
MODES = {"D": "decomposed", "A": "arbitrated"}
N_BOOT, SEED = 2000, 0


def fnum(v):
    try:
        return float(v)
    except (TypeError, ValueError):
        return None


def load_rows(xlsx: Path) -> dict[int, dict]:
    """student -> {ai, ta, t1, t2, q:{k:(ai, ta)}} for rows with complete Q1-Q3 data."""
    ws = openpyxl.load_workbook(xlsx, data_only=True).active
    hdr = {c.value: i for i, c in enumerate(ws[1]) if c.value}
    out = {}
    for r in ws.iter_rows(min_row=2, values_only=True):
        if r[0] is None:
            continue
        q, ok, t1s, t2s = {}, True, 0.0, 0.0
        for k in (1, 2, 3):
            ai = fnum(r[hdr[f"AI Q{k} Score"]])
            t1, t2 = fnum(r[hdr[f"TA 1 - Q{k} Score"]]), fnum(r[hdr[f"TA 2 - Q{k} Score"]])
            if ai is None or t1 is None or t2 is None:
                ok = False
                break
            q[k] = (ai, (t1 + t2) / 2)
            t1s += t1
            t2s += t2
        if not ok:
            continue
        out[int(r[0])] = {"ai": sum(v[0] for v in q.values()), "ta": (t1s + t2s) / 2,
                          "t1": t1s, "t2": t2s, "q": q}
    return out


def boot_ci(vals: list[float]) -> tuple[float, float]:
    rng = np.random.default_rng(SEED)
    a = np.asarray(vals)
    idx = rng.integers(0, len(a), size=(N_BOOT, len(a)))
    means = np.abs(a[idx]).mean(axis=1)
    return float(np.percentile(means, 2.5)), float(np.percentile(means, 97.5))


def stats(rows: dict[int, dict], students: list[int]) -> dict:
    xs = [rows[s] for s in students if s in rows]
    n = len(xs)
    if n == 0:
        return {"n": 0}
    res = [x["ai"] - x["ta"] for x in xs]
    tot = [x["ai"] for x in xs]
    mae = statistics.fmean(abs(e) for e in res)
    lo, hi = boot_ci(res)
    sd = statistics.pstdev(tot) if n > 1 else 0.0
    zero = sum(1 for t in tot if t == 0) / n
    beh = ("refusal" if zero >= 0.90 and sd < 0.5 else "near-refusal" if zero >= 0.90
           else "collapse" if mae >= 8 else "graded")
    r = float(np.corrcoef(tot, [x["ta"] for x in xs])[0, 1]) if sd > 0 else float("nan")
    return {"n": n, "mae": mae, "lo": lo, "hi": hi, "bias": statistics.fmean(res), "zero": zero,
            "sd": sd, "mean": statistics.fmean(tot), "r": r, "beh": beh,
            "q": {k: statistics.fmean(abs(x["q"][k][0] - x["q"][k][1]) for x in xs) for k in (1, 2, 3)},
            "floor": statistics.fmean(abs(x["t1"] - x["t2"]) for x in xs)}


def paper_run(rid: str) -> Path:
    hits = sorted(PAPER_RESULTS.glob(f"{rid}__*.xlsx"))
    if not hits:
        sys.exit(f"no paper workbook for {rid}")
    return hits[0]


def fmt(label: str, s: dict) -> str:
    if s["n"] == 0:
        return f"    {label:<34}  (no overlapping students)"
    return (f"    {label:<34} n={s['n']:>3}  MAE {s['mae']:6.3f} [{s['lo']:.3f}, {s['hi']:.3f}]  bias {s['bias']:+7.3f}"
            f"  zero {100*s['zero']:5.1f}%  sd {s['sd']:5.2f}  mean {s['mean']:5.2f}  r {s['r']:5.2f}  {s['beh']:<12}"
            f"  Q1/Q2/Q3 MAE {s['q'][1]:.2f}/{s['q'][2]:.2f}/{s['q'][3]:.2f}")


def main() -> None:
    ap = argparse.ArgumentParser()
    ap.add_argument("--models", nargs="*", default=list(MODELS))
    a = ap.parse_args()
    print("# Prompt-only baselines under the strict persona (CV exam, students 1-100, Q1-Q3 base scale)")
    print(f"# numpy {np.__version__}, openpyxl {openpyxl.__version__}; bootstrap {N_BOOT} resamples, default_rng({SEED})")
    print("# Every paper run is restricted to the students the baseline graded (paired comparison).\n")
    summary = []
    for idx in a.models:
        name, strict_rid, neutral_rid = MODELS[idx]
        paper_strict = load_rows(paper_run(strict_rid))
        paper_neutral = load_rows(paper_run(neutral_rid))
        print(f"## {name}  (paper strict {strict_rid}, neutral {neutral_rid})")
        for mcode, mode in MODES.items():
            hits = sorted(PB_RESULTS.glob(f"PB-{mcode}{idx}__*.xlsx"))
            if not hits:
                print(f"    PB-{mcode}{idx} ({mode}): no workbook yet")
                continue
            rows = load_rows(hits[0])
            students = sorted(rows)
            if not students:
                print(f"    PB-{mcode}{idx} ({mode}): workbook has no complete rows yet")
                continue
            print(f"  {hits[0].name}")
            print(f"    students with complete Q1-Q3: {len(students)} ({students[0]}..{students[-1]}); "
                  f"human floor on them {stats(rows, students)['floor']:.3f}")
            s_new, s_str, s_neu = stats(rows, students), stats(paper_strict, students), stats(paper_neutral, students)
            print(fmt(f"{mode} (strict persona)", s_new))
            print(fmt(f"paper strict {strict_rid}", s_str))
            print(fmt(f"paper neutral {neutral_rid}", s_neu))
            if s_str["n"] and s_neu["n"]:
                print(f"    -> MAE vs paper strict {s_new['mae'] - s_str['mae']:+.3f}; vs paper neutral "
                      f"{s_new['mae'] - s_neu['mae']:+.3f}; ratio to neutral x{s_new['mae'] / s_neu['mae']:.2f} "
                      f"(paper strict x{s_str['mae'] / s_neu['mae']:.2f}); in band: {s_new['beh'] == 'graded'}")
                summary.append((name, mode, s_new, s_str, s_neu))
        print()
    if summary:
        print("## Summary (MAE on the paired students; band = MAE < 8)")
        print(f"    {'model':<20} {'mode':<12} {'n':>3} {'baseline':>9} {'paper strict':>13} {'paper neutral':>14}  verdict")
        for name, mode, s_new, s_str, s_neu in summary:
            verdict = ("repairs to neutral level" if s_new["mae"] <= s_neu["mae"] + 0.5 else
                       "in band, above neutral" if s_new["beh"] == "graded" else f"still {s_new['beh']}")
            print(f"    {name:<20} {mode:<12} {s_new['n']:>3} {s_new['mae']:>9.3f} {s_str['mae']:>13.3f} "
                  f"{s_neu['mae']:>14.3f}  {verdict}")


if __name__ == "__main__":
    main()
