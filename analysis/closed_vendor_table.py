#!/usr/bin/env python3
"""closed_vendor_table.py -- the cross-vendor closed-model persona table: Gemini (paper runs + F03/IG23), OpenAI and
Anthropic (runs O*/N*/IO*/IN*, graded via closed_batch_grade.py and part of the
grids since 2026-09-21) under neutral / strict / lenient on both exams. Numbers
are read from the two master comparison CSVs, whose conventions are:

  * CV exam: AI base total (Q1-Q3, 35 points) vs the grader-average base total.
  * ML exam: AI score+bonus vs the graders' total grades (~65 points).
  * MAE with a 2,000-resample percentile bootstrap CI, bias = mean signed error,
    Pearson r, zero-total rate, and the Section 5.1 behaviour class
    (refusal: >=90% zero totals and sd<0.5; collapse: MAE >= 3.07x the exam's floor,
    i.e. 8 on the CV exam and 15.7 on the ML exam; graded otherwise).

Writes analysis/closed_vendor_table.md and analysis/closed_vendor_table.tex.
Run:  python analysis/closed_vendor_table.py
"""
from __future__ import annotations

import glob
import json
import statistics
from pathlib import Path

import numpy as np
import openpyxl

ROOT = Path(__file__).resolve().parent.parent
CV_RES = ROOT / "computer_vision_results" / "results"
ML_RES = ROOT / "introduction_to_ai_results" / "results"
FLOOR = {"cv": 2.61, "ml": 5.13}
N_BOOT = 2000

# Paper runs (Gemini) to include, by exam: run id -> (model label, persona)
GEMINI = {
    "cv": {"D02": ("gemini-3.1-pro-preview", "neutral"), "F03": ("gemini-3.1-pro-preview", "strict"),
           "F02": ("gemini-3.1-pro-preview", "lenient"),
           "A01": ("gemini-flash-lite", "neutral"), "C01": ("gemini-flash-lite", "strict"),
           "C02": ("gemini-flash-lite", "lenient")},
    "ml": {"IG08": ("gemini-3.1-pro-preview", "neutral"), "IG23": ("gemini-3.1-pro-preview", "strict"),
           "IG14": ("gemini-3.1-pro-preview", "lenient"),
           "IG01": ("gemini-flash-lite", "neutral"), "IG05": ("gemini-flash-lite", "strict"),
           "IG06": ("gemini-flash-lite", "lenient")},
}
VENDOR = {"gemini": "Google", "gpt": "OpenAI", "claude": "Anthropic"}
PERSONAS = ["neutral", "strict", "lenient"]


def vendor(model: str) -> str:
    for k, v in VENDOR.items():
        if model.startswith(k):
            return v
    return "?"


def stats(xlsx: Path, exam: str) -> dict | None:
    ws = openpyxl.load_workbook(xlsx, data_only=True).active
    hdr = [c.value for c in ws[1]]
    col = {h: i for i, h in enumerate(hdr) if h}
    if exam == "cv":
        t1, t2 = col["TA 1 - Total Score (out of 35)"], col["TA 2 - Total Score (out of 35)"]
        ai_cols = [col["AI Total Score"]]
    else:
        t1, t2 = col["TA 1 - Total Grade"], col["TA 2 - Total Grade"]
        ai_cols = [col["AI Total Score"], col["AI Total Bonus"]]
    ai, ta, q1zero = [], [], 0
    for r in ws.iter_rows(min_row=2, values_only=True):
        vals = [r[c] for c in ai_cols]
        if any(v in (None, "") for v in vals) or r[t1] in (None, "") or r[t2] in (None, ""):
            continue
        ai.append(sum(float(v) for v in vals))
        ta.append((float(r[t1]) + float(r[t2])) / 2)
        q1zero += (r[col["AI Q1 Score"]] == 0)
    n = len(ai)
    if n < 2:
        return None
    d = [a - t for a, t in zip(ai, ta)]
    mae = statistics.fmean(abs(x) for x in d)
    # The master CSVs' bootstrap (run_analysis.bootstrap_mae_ci): numpy
    # default_rng(0), percentile. A random.Random(0) resample here gave the
    # prompt-baseline rows of tab:allruns CIs that differed in the second
    # decimal from the same runs in tab:promptbaselines.
    rng = np.random.default_rng(0)
    idx = rng.integers(0, n, size=(N_BOOT, n))
    lo, hi = np.percentile(np.abs(np.asarray(d))[idx].mean(axis=1), [2.5, 97.5])
    zero_rate = sum(a == 0 for a in ai) / n
    sd = statistics.stdev(ai)
    if zero_rate >= 0.9 and sd < 0.5:
        beh = "refusal"
    elif zero_rate >= 0.9:
        beh = "near-refusal"
    elif mae >= (8.0 if exam == "cv" else 15.7):   # floor-matched band: 3.07x each exam's floor (item 8)
        beh = "collapse"
    else:
        beh = "graded"
    return {"n": n, "mae": mae, "lo": float(lo), "hi": float(hi),
            "bias": statistics.fmean(d), "r": statistics.correlation(ai, ta) if sd > 0 else float("nan"), "zero_rate": zero_rate,
            "mean_awarded": statistics.fmean(ai), "sd": sd, "q1zero": q1zero, "behaviour": beh}


def collect() -> dict[tuple[str, str, str], dict]:
    """(model, exam, persona) -> stats (+ run id, cost), read from the two master
    CSVs so this table prints exactly the numbers of the appendix run tables
    (the closed runs entered the grids on 2026-09-21). Cost comes from
    closed_runs/<RID>/meta.json; stats() above is kept for callers that need
    the same conventions on a workbook outside the CSVs."""
    import csv
    cost = {}
    for meta_p in glob.glob(str(ROOT / "closed_runs" / "*" / "meta.json")):
        meta = json.loads(Path(meta_p).read_text())
        cost[meta["run_id"]] = (meta.get("cost") or {}).get("usd_batch")
    out: dict[tuple[str, str, str], dict] = {}
    for exam, csv_path in (("cv", ROOT / "analysis" / "computer_vision_master_comparison.csv"),
                           ("ml", ROOT / "analysis" / "introduction_to_ai_master_comparison.csv")):
        for r in csv.DictReader(open(csv_path, encoding="utf-8")):
            rid = r["run_id"]
            if rid in GEMINI[exam]:
                model, persona = GEMINI[exam][rid]
            elif r["source"] in ("openai", "anthropic"):
                model, persona = r["model"], r["strictness"]
            else:
                continue
            f = lambda k: float(r[k]) if r.get(k) not in (None, "", "nan") else float("nan")
            out[(model, exam, persona)] = {
                "n": int(float(r["n_valid_vs_TA"])), "mae": f("MAE_vs_TA_avg"), "lo": f("MAE_CI95_lo"),
                "hi": f("MAE_CI95_hi"), "bias": f("bias_vs_TA_avg"), "r": f("AI_TA_pearson_r"),
                "zero_rate": f("zero_rate"), "behaviour": str(r["behaviour"]).split(" (")[0],
                "run_id": rid, "cost": cost.get(rid)}
    return out


def fmt(s: dict | None, floor: float) -> str:
    if not s:
        return "—"
    return f"{s['mae']:.2f} [{s['lo']:.2f}, {s['hi']:.2f}] {s['bias']:+.2f}"


def main() -> None:
    data = collect()
    models = sorted({k[0] for k in data}, key=lambda m: (vendor(m), m))
    md = ["# Closed models across vendors under neutral / strict / lenient",
          "", "MAE [95% CI] bias, against the grader average; CV floor 2.61/35, ML floor 5.13/65. "
          "Behaviour per Section 5.1. Gemini rows are the paper's runs (plus F03/IG23); OpenAI and Anthropic rows come from "
          "closed_runs/ (Batch API, temperature 0 where the API allows it, reasoning/thinking off).", ""]
    md.append("| vendor | model | persona | CV: MAE [CI] bias | CV n | CV behaviour | ML: MAE [CI] bias | ML n | ML behaviour | run ids | batch cost |")
    md.append("|---|---|---|---|---|---|---|---|---|---|---|")
    tex = ["% Auto-generated by analysis/closed_vendor_table.py -- do not edit by hand.",
           r"\begin{tabular}{llrlrl}", r"\toprule",
           r"Model & Persona & \multicolumn{2}{c}{CV exam (floor $2.61$)} & \multicolumn{2}{c}{ML exam (floor $5.13$)} \\",
           r"\cmidrule(lr){3-4}\cmidrule(lr){5-6}",
           r" & & MAE [95\% CI] & bias & MAE [95\% CI] & bias \\", r"\midrule"]
    for m in models:
        for i, p in enumerate(PERSONAS):
            cv, ml = data.get((m, "cv", p)), data.get((m, "ml", p))
            if not cv and not ml:
                continue
            ids = "/".join(x["run_id"] for x in (cv, ml) if x)
            cost = sum(x["cost"] for x in (cv, ml) if x and x["cost"]) or None
            md.append(f"| {vendor(m)} | {m} | {p} | {fmt(cv, FLOOR['cv'])} | {cv['n'] if cv else '—'} | "
                      f"{cv['behaviour'] if cv else '—'} | {fmt(ml, FLOOR['ml'])} | {ml['n'] if ml else '—'} | "
                      f"{ml['behaviour'] if ml else '—'} | {ids} | {('$%.2f' % cost) if cost else '—'} |")
            label = m if i == 0 else ""
            def cell(s):
                if not s:
                    return "--- & ---"
                mae = f"$\\mathbf{{{s['mae']:.2f}}}$" if s["behaviour"] != "graded" else f"${s['mae']:.2f}$"
                return f"{mae} $[{s['lo']:.2f},\\,{s['hi']:.2f}]$ & ${s['bias']:+.2f}$"
            tex.append(f"{label} & \\emph{{{p}}} & {cell(cv)} & {cell(ml)} \\\\")
        tex.append(r"\midrule")
    tex[-1] = r"\bottomrule"
    tex.append(r"\end{tabular}")
    md += ["", "Ratios strict/neutral and lenient/neutral (MAE):", ""]
    for m in models:
        n_cv, n_ml = data.get((m, "cv", "neutral")), data.get((m, "ml", "neutral"))
        parts = []
        for p in ("strict", "lenient"):
            for exam, base in (("cv", n_cv), ("ml", n_ml)):
                s = data.get((m, exam, p))
                if s and base:
                    parts.append(f"{p}/{exam.upper()} ×{s['mae'] / base['mae']:.2f} (zero-rate {s['zero_rate']:.1%}, r={s['r']:.3f})")
        if parts:
            md.append(f"- {m}: " + "; ".join(parts))
    (ROOT / "analysis" / "closed_vendor_table.md").write_text("\n".join(md) + "\n")
    (ROOT / "analysis" / "closed_vendor_table.tex").write_text("\n".join(tex) + "\n")
    (ROOT / "sections").mkdir(exist_ok=True)   # sections/ does not ship with the code
    (ROOT / "sections" / "closed_vendors_table.tex").write_text("\n".join(tex) + "\n")   # \input by Appendix C
    print("\n".join(md))
    print("\nwrote analysis/closed_vendor_table.md and .tex")


if __name__ == "__main__":
    main()
