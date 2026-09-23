#!/usr/bin/env python3
"""Emit the LaTeX bodies for the second-exam paper tables, plus every derived
statistic the replication prose quotes -- the Stage 2 counterpart of
computer_vision_make_paper_tables.py.

Why a generator: Table tab:replication went from 9 hand-typed model rows to 17
across seven families, and the surrounding prose quotes a dozen numbers no
single report file carries (the calibration correlation, the Fisher z against
Stage 3, the matched few-shot subsample). Everything here is derived from
analysis/introduction_to_ai_master_comparison.csv, the result workbooks and the
ground truth; nothing is typed.

The team's convention is hand-written LaTeX in sections/*.tex, so this prints
to stdout for pasting rather than writing \\input files. Rerun after any
pipeline change and diff against the .tex to catch drift:

    python analysis/introduction_to_ai_make_paper_tables.py
"""
import sys
from pathlib import Path

import numpy as np
import pandas as pd

sys.path.insert(0, str(Path(__file__).parent))
import introduction_to_ai_run_analysis as ra  # noqa: E402

BASE = Path(__file__).resolve().parent.parent
CSV = Path(__file__).parent / "introduction_to_ai_master_comparison.csv"
CV_CSV = Path(__file__).parent / "computer_vision_master_comparison.csv"
RESULTS = BASE / "introduction_to_ai_results" / "results"

# Display names, keyed on the model tag so a new model fails loudly rather
# than silently printing a tag where a name belongs.
DISPLAY = {
    "qwen25-coder-7b": "Qwen2.5-Coder-7B", "qwen25-coder-14b": "Qwen2.5-Coder-14B",
    "qwen25-coder-32b": "Qwen2.5-Coder-32B", "qwen25-72b": "Qwen2.5-72B",
    "qwen3-coder-30b-a3b": "Qwen3-Coder-30B-A3B", "qwen3-coder-next": "Qwen3-Coder-Next 80B",
    "qwen3-235b-a22b": "Qwen3-235B-A22B", "qwen3-coder-480b": "Qwen3-Coder-480B",
    "llama31-8b": "Llama-3.1-8B", "llama33-70b": "Llama-3.3-70B",
    "glm4-9b": "GLM-4-9B", "glm4-32b": "GLM-4-32B", "glm45-air": "GLM-4.5-Air",
    "gemma3-12b": "Gemma-3-12B", "gemma3-27b": "Gemma-3-27B",
    "mistral-small-24b": "Mistral-Small-24B",
    "deepseek-coder-v2-lite": "DeepSeek-Coder-V2-Lite",
    "flash-lite": "Flash-Lite", "3-flash-preview": "3-flash-preview",
    "3.1-pro-preview": "3.1-pro-preview", "2.5-pro": "2.5-pro",
}


def default_pairs(d):
    """Matched neutral/strict open-weights pairs at the default prompt config."""
    sel = d[(d.source == "local") & (d.use_solution == 1) & (d.use_breakdown == 1)
            & (d.use_reasoning == 0) & (d.few_shot == 0) & (d.temperature == 0)
            & (d.runs == 1) & (d.students == "all")]
    piv = {}
    for _, r in sel.iterrows():
        piv.setdefault((r.model, r.strictness), r)
    out = []
    for m in sorted(sel.model.unique()):
        n, s = piv.get((m, "neutral")), piv.get((m, "strict"))
        if n is None or s is None:
            continue
        out.append((m, n, s))
    return out


def floor_stats():
    rows = ra.ta_rows()
    gaps = [abs(r["t1"] - r["t2"]) for r in rows]
    return float(np.mean(gaps))


def behaviour_tex(b):
    b = str(b).split(" (")[0]
    return "\\textbf{refusal}" if b == "refusal" else b


def table_replication(pairs, floor):
    print("% ---- tab:replication body (17 matched pairs, sorted by damage ratio) ----")
    rows = sorted(pairs, key=lambda t: -(t[2].MAE_vs_TA_avg / t[1].MAE_vs_TA_avg))
    for m, n, s in rows:
        ratio = s.MAE_vs_TA_avg / n.MAE_vs_TA_avg
        print(f"{DISPLAY[m]} & {n.MAE_vs_TA_avg:.2f} & {s.MAE_vs_TA_avg:.2f} & "
              f"[{s.MAE_CI95_lo:.2f}, {s.MAE_CI95_hi:.2f}] & $\\times {ratio:.2f}$ & "
              f"{s.MAE_vs_TA_avg / floor:.1f} & {s.zero_rate:.1%} & "
              f"${s.bias_vs_TA_avg:+.2f}$ & {behaviour_tex(s.behaviour)} \\\\"
              .replace("%", "\\%"))
    print()


def table_repmech(d):
    """Blanket vs selective zeroing on the strict runs, from the CSV's Q1/Q2 means
    is not enough -- the counts come from the ledger's run stats, so recompute.

    Two gates, both from ra: the run must have left the graded band, and it must
    have zeroed Q1 on at least ra.MIN_Q1_ZERO_RATE of the students it graded.
    Below that volume the zeros are ordinary non-attempts and the selective
    share is noise on a handful of students (item 60)."""
    TA = ra.ta_truth()
    print("% ---- tab:repmech body (Q1-zero probe on the 17 strict runs) ----")
    sel = []
    for m, n, s in default_pairs(d):
        stats = ra.run_stats(find_workbook(s.run_id), TA)
        share = stats["q1zero_q2pos"] / stats["q1zero"] if stats["q1zero"] else 0.0
        rate = stats["q1zero"] / stats["n"]        # stats["n"] = students graded
        beh = str(s.behaviour).split(" (")[0]
        if beh not in ("collapse", "refusal", "near-refusal"):
            reading = "n/a (in band)"
        elif rate < ra.MIN_Q1_ZERO_RATE:
            reading = "n/a (too few Q$1$ zeros)"
        else:
            reading = "selective" if share >= ra.MECH_SELECTIVE else "blanket"
        sel.append((m, stats["q1zero"], stats["n"], rate, stats["q1zero_q2pos"],
                    share, reading))
    for m, q1z, n_, rate, q2p, share, reading in sorted(sel, key=lambda t: -t[1]):
        print(f"{DISPLAY[m]} & {q1z} / {n_} ({rate:.0%}) & {q2p} & {share:.0%} & "
              f"{reading} \\\\".replace("%", "\\%"))
    print()


_WB_BY_RUN = None


def find_workbook(run_id):
    """The earliest runs carry a run_id that differs from the filename stem,
    so resolve through the tracker's output_xlsx column, as all_runs() does."""
    global _WB_BY_RUN
    if _WB_BY_RUN is None:
        import openpyxl
        ws = openpyxl.load_workbook(BASE / "introduction_to_ai_results" / "ablation_runs.xlsx",
                                    data_only=True).active
        hdr = [c.value for c in ws[1]]
        c = {h: i for i, h in enumerate(hdr)}
        _WB_BY_RUN = {str(r[c["run_id"]]): str(r[c["output_xlsx"]] or "")
                      for r in ws.iter_rows(min_row=2, values_only=True) if r[c["run_id"]]}
    rel = _WB_BY_RUN.get(run_id, "")
    path = BASE / "introduction_to_ai_results" / rel if rel else None
    if path is None or not path.exists():
        hits = list(RESULTS.glob(f"{run_id}__*.xlsx"))
        assert len(hits) == 1, (run_id, hits)
        return hits[0]
    return path


def calibration_stats(pairs, d_cv):
    """Neutral bias vs strict effect: the calibration-correction correlation."""
    rows = [(m, n.bias_vs_TA_avg, s.MAE_vs_TA_avg - n.MAE_vs_TA_avg,
             str(s.behaviour).startswith("refusal")) for m, n, s in pairs]
    print("% ---- calibration correlation (neutral bias vs strict effect) ----")
    for label, sub in (("all", rows), ("excl. refusal", [r for r in rows if not r[3]])):
        b = np.array([r[1] for r in sub])
        e = np.array([r[2] for r in sub])
        r_p = float(np.corrcoef(b, e)[0, 1])
        r_s = float(pd.Series(b).rank().corr(pd.Series(e).rank()))
        p = ra.perm_p(list(b), list(e), lambda x, y: abs(ra.pearson(x, y)))
        print(f"%   Stage 2 {label} (n={len(sub)}): pearson {r_p:+.3f} "
              f"spearman {r_s:+.3f} perm_p {p:.4f}")
    improved = sum(1 for r in rows if r[2] < 0)
    print(f"%   improved under strict: {improved} of {len(rows)}")

    # Stage 3 counterpart at its default config, for the contrast.
    cv = d_cv[(d_cv.source == "local") & (d_cv.use_solution == 1)
              & (d_cv.use_guidelines == 1) & (d_cv.use_breakdown == 1)
              & (d_cv.use_reasoning == 0) & (d_cv.few_shot == 0)
              & (d_cv.temperature == 0) & (d_cv.runs == 1)
              & (d_cv.n_valid_vs_TA >= 560)]
    piv = {}
    for _, r in cv.iterrows():
        piv.setdefault((r.model, r.strictness), r)
    cv_rows = []
    for m in sorted(cv.model.unique()):
        n, s = piv.get((m, "neutral")), piv.get((m, "strict"))
        if n is None or s is None:
            continue
        cv_rows.append((n.bias_vs_TA_avg, s.MAE_vs_TA_avg - n.MAE_vs_TA_avg,
                        str(s.behaviour) == "refusal"))
    for label, sub in (("all", cv_rows), ("excl. refusal", [r for r in cv_rows if not r[2]])):
        b = np.array([r[0] for r in sub])
        e = np.array([r[1] for r in sub])
        print(f"%   Stage 3 {label} (n={len(sub)}): pearson "
              f"{float(np.corrcoef(b, e)[0, 1]):+.3f}")
    print("%   NB the Stage 3 r flips sign if the near-refusal's ceiling MAE is"
          " also dropped -- its biases barely vary, so no r there is stable;"
          " quote '0 of 16 improve', never a Stage 3 correlation.")
    print(f"%   Stage 3 improved under strict: "
          f"{sum(1 for r in cv_rows if r[1] < 0)} of {len(cv_rows)}")
    print()


def pair_inversion_stats(d):
    """Per-pair human noise vs AI error, IG08 side, plus the Fisher z vs Stage 3."""
    rows = ra.ta_rows()
    stable, _ = ra.pair_groups(rows)
    TA = ra.ta_truth()
    wb = find_workbook("IG08")
    import openpyxl
    ws = openpyxl.load_workbook(wb, data_only=True).active
    hdr = [c.value for c in ws[1]]
    c = {h: i for i, h in enumerate(hdr)}
    ai = {}
    for r in ws.iter_rows(min_row=2, values_only=True):
        a = r[c["AI Total Score"]]
        if a is None:
            continue
        ai[r[c["Number"]]] = float(a) + float(r[c["AI Total Bonus"]] or 0)
    pts, corr = [], []
    for pair, members in stable.items():
        h = float(np.mean([abs(r["t1"] - r["t2"]) for r in members]))
        errs = [ai[r["n"]] - (r["t1"] + r["t2"]) / 2 for r in members if r["n"] in ai]
        gaps = [r["t1"] - r["t2"] for r in members if r["n"] in ai]
        mse = float(np.mean(np.square(errs)))
        inj = float(np.var(gaps)) / 4.0
        pts.append((h, float(np.mean(np.abs(errs)))))
        corr.append((h, max(mse - inj, 0.0) ** 0.5))
    pts.sort()  # same point order as the pair-level report, so the seeded
    hx = [p_[0] for p_ in pts]  # permutation p reproduces to the digit
    ay = [p_[1] for p_ in pts]
    r_p = ra.pearson(hx, ay)
    r_s = ra.spearman(hx, ay)
    p = ra.perm_p(hx, ay, lambda x, y: abs(ra.pearson(x, y)))
    z, pz = ra.fisher_z(r_p, len(hx), ra.S3["gt_r"], ra.S3["gt_k"])
    print("% ---- per-pair inversion, AI side = IG08 (best full-cohort run) ----")
    print(f"%   pearson {r_p:+.3f} spearman {r_s:+.3f} perm_p {p:.4f} over {len(hx)} pairs")
    print(f"%   AI per-pair span x{max(ay)/min(ay):.2f}, human span x{max(hx)/min(hx):.2f}")
    print(f"%   Fisher z vs Stage 3 D01 ({ra.S3['gt_r']:+.3f}, {ra.S3['gt_k']} pairs): "
          f"z = {z:.2f}, p = {pz:.3f}")
    corr.sort()
    r_c = ra.pearson([c[0] for c in corr], [c[1] for c in corr])
    print(f"%   after subtracting the injected Var(TA1-TA2)/4 from per-pair MSE: r = {r_c:+.3f}")
    print()


def fewshot_stats(d):
    """The two few-shot arms; IG18 must be quoted on the matched subsample,
    because its dropout is biased toward stronger students."""
    import openpyxl
    from scipy import stats as st

    def totals(run_id):
        ws = openpyxl.load_workbook(find_workbook(run_id), data_only=True).active
        hdr = [c.value for c in ws[1]]
        c = {h: i for i, h in enumerate(hdr)}
        out = {}
        for r in ws.iter_rows(min_row=2, values_only=True):
            a = r[c["AI Total Score"]]
            if a is None:
                continue
            out[r[c["Number"]]] = float(a) + float(r[c["AI Total Bonus"]] or 0)
        return out

    TA = {n: (a + b) / 2 for n, (a, b) in ra.ta_truth().items()
          if a is not None and b is not None}
    print("% ---- few-shot ----")
    a201 = d.set_index("run_id")
    for rid, base in (("IA201", "IA01-Q30-n"),):
        print(f"%   {base}: MAE {a201.loc[base].MAE_vs_TA_avg:.2f} bias "
              f"{a201.loc[base].bias_vs_TA_avg:+.2f}  ->  {rid}: MAE "
              f"{a201.loc[rid].MAE_vs_TA_avg:.2f} bias {a201.loc[rid].bias_vs_TA_avg:+.2f} "
              f"(both full cohort)")
    base, few = totals("IG01"), totals("IG18")
    matched = sorted(set(base) & set(few) & set(TA))
    dropped = sorted((set(base) - set(few)) & set(TA))
    bm = np.array([base[n] - TA[n] for n in matched])
    fm = np.array([few[n] - TA[n] for n in matched])
    print(f"%   IG18 matched subsample (n={len(matched)}): IG01 MAE "
          f"{np.abs(bm).mean():.2f} bias {bm.mean():+.2f}  ->  IG18 MAE "
          f"{np.abs(fm).mean():.2f} bias {fm.mean():+.2f}")
    t, p = st.ttest_ind([TA[n] for n in dropped], [TA[n] for n in matched],
                        equal_var=False)
    print(f"%   dropout bias: dropped n={len(dropped)} TA mean "
          f"{np.mean([TA[n] for n in dropped]):.2f} vs kept "
          f"{np.mean([TA[n] for n in matched]):.2f} (Welch p = {p:.1e})")
    print()


def subset_stats():
    """gemini-3-flash-preview subset runs (IG15--IG20) need a baseline on the
    same first-100 students; the full-cohort IG07 over-credits them."""
    import openpyxl
    ws = openpyxl.load_workbook(find_workbook("IG07"), data_only=True).active
    hdr = [c.value for c in ws[1]]
    c = {h: i for i, h in enumerate(hdr)}
    diffs = []
    for r in ws.iter_rows(min_row=2, values_only=True):
        n, a = r[c["Number"]], r[c["AI Total Score"]]
        if a is None or n > 100:
            continue
        t1, t2 = r[c["TA 1 - Total Grade"]], r[c["TA 2 - Total Grade"]]
        diffs.append(float(a) + float(r[c["AI Total Bonus"]] or 0) - (t1 + t2) / 2)
    mae = float(np.mean(np.abs(diffs)))
    lo, hi = ra.boot_ci([abs(x) for x in diffs])
    print("% ---- IG07 restricted to the first-100 subset ----")
    print(f"%   n={len(diffs)}: MAE {mae:.2f} [{lo:.2f}, {hi:.2f}] "
          f"bias {float(np.mean(diffs)):+.2f}")
    print()


def cross_exam_neutral(d, d_cv):
    """Neutral MAE correlation across the 17 models run on both exams."""
    KEY = {"qwen25-coder-7b": "qwen2.5-coder-7b", "qwen25-coder-14b": "qwen2.5-coder-14b",
           "qwen25-coder-32b": "qwen2.5-coder-32b", "qwen25-72b": "qwen2.5-72b",
           "qwen3-coder-30b-a3b": "qwen3-coder-30b-a3b", "qwen3-coder-next": "qwen3-coder-next",
           "qwen3-235b-a22b": "qwen3-235b-a22b", "qwen3-coder-480b": "qwen3-coder-480b",
           "llama31-8b": "llama-3.1-8b", "llama33-70b": "llama-3.3-70b",
           "glm4-9b": "glm-4-9b", "glm4-32b": "glm-4-32b", "glm45-air": "glm-4.5-air",
           "gemma3-12b": "gemma-3-12b", "gemma3-27b": "gemma-3-27b",
           "mistral-small-24b": "mistral-small-24b",
           "deepseek-coder-v2-lite": "deepseek-coder-v2-lite"}
    ia = {KEY[m]: n.MAE_vs_TA_avg for m, n, _ in default_pairs(d)}
    cv = d_cv[(d_cv.source == "local") & (d_cv.use_solution == 1)
              & (d_cv.use_guidelines == 1) & (d_cv.use_breakdown == 1)
              & (d_cv.use_reasoning == 0) & (d_cv.few_shot == 0)
              & (d_cv.temperature == 0) & (d_cv.runs == 1)
              & (d_cv.n_valid_vs_TA >= 560) & (d_cv.strictness == "neutral")]
    cvm = cv.drop_duplicates("model").set_index("model").MAE_vs_TA_avg
    common = sorted(set(ia) & set(cvm.index))
    x = [cvm[m] for m in common]
    y = [ia[m] for m in common]
    print("% ---- cross-exam neutral MAE ----")
    print(f"%   over {len(common)} models: pearson {ra.pearson(x, y):+.3f} "
          f"spearman {ra.spearman(x, y):+.3f}")
    print()


def main():
    d = pd.read_csv(CSV)
    d_cv = pd.read_csv(CV_CSV)
    floor = floor_stats()
    pairs = default_pairs(d)
    print(f"% floor = {floor:.3f}; matched pairs = {len(pairs)}\n")
    table_replication(pairs, floor)
    table_repmech(d)
    calibration_stats(pairs, d_cv)
    pair_inversion_stats(d)
    fewshot_stats(d)
    subset_stats()
    cross_exam_neutral(d, d_cv)


if __name__ == "__main__":
    main()
