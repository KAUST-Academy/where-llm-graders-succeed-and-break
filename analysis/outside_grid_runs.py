#!/usr/bin/env python3
"""Rows for the prompt-only baselines of Appendix B (prompt_baselines/results/,
written by prompt_baselines_grade.py; Table tab:promptbaselines) in the appendix
run-table schema (run id, model, persona, prompt flags, t, n, MAE, 95% CI, bias,
behaviour). They are graded with a different instrument (one call per rubric
row, or strict + neutral + arbiter), so they are listed after the grid and not
counted as grader configurations. closed_rows() is kept for the cross-vendor
closed runs' provenance (closed_runs/<RID>/meta.json); those runs themselves
now live in the results directories and enter the master CSVs like any other.

Metrics follow the paper's conventions (Section 3) through
closed_vendor_table.stats: CV exam = AI base total vs the grader-average base
total on the 35-point scale; ML exam = AI score+bonus vs the graders' totals;
MAE with a 2,000-resample percentile bootstrap CI; bias = mean signed error;
behaviour classes at each exam's floor-matched band. The prompt-baseline
workbooks store the graders' totals as uncached formulas, so those totals are
rebuilt from the per-question columns before the same statistics are taken.

Used by computer_vision_make_appendix_table.py and
introduction_to_ai_make_appendix_table.py; run directly to print the rows.
"""
from __future__ import annotations

import glob
import json
import re
import sys
import tempfile
from pathlib import Path

import openpyxl

sys.path.insert(0, str(Path(__file__).resolve().parent))
from closed_vendor_table import stats  # noqa: E402  (same conventions as Table tab:closedvendors)

ROOT = Path(__file__).resolve().parent.parent
CLOSED = ROOT / "closed_runs"
PB = ROOT / "prompt_baselines" / "results"


def prompt_flags_from_config(config: str) -> str:
    """'sol1_gd1_rs0_bd1' -> 'SGB' (S solution, G guidelines, B breakdown, R thinking)."""
    tok = dict(re.findall(r"(sol|gd|rs|bd)(\d)", config))
    s = ("S" if tok.get("sol") == "1" else "") + ("G" if tok.get("gd") == "1" else "") \
        + ("B" if tok.get("bd") == "1" else "") + ("R" if tok.get("rs") == "1" else "")
    return s or "---"


def _with_rebuilt_totals(xlsx: Path) -> Path:
    """Copy a CV workbook with the two grader totals filled from Q1-Q3 where the
    stored total is an uncached formula (prompt_baselines_grade.py output)."""
    wb = openpyxl.load_workbook(xlsx)
    ws = wb.active
    hdr = [c.value for c in ws[1]]
    col = {h: i + 1 for i, h in enumerate(hdr) if h}
    for ta in (1, 2):
        tot = col[f"TA {ta} - Total Score (out of 35)"]
        qs = [col[f"TA {ta} - Q{k} Score"] for k in (1, 2, 3)]
        for r in range(2, ws.max_row + 1):
            v = ws.cell(r, tot).value
            if isinstance(v, (int, float)):
                continue
            parts = [ws.cell(r, q).value for q in qs]
            ws.cell(r, tot).value = (sum(float(p) for p in parts)
                                     if all(isinstance(p, (int, float)) for p in parts) else None)
    tmp = Path(tempfile.mkdtemp()) / xlsx.name
    wb.save(tmp)
    return tmp


def closed_rows(exam: str) -> list[dict]:
    rows = []
    for meta_p in sorted(glob.glob(str(CLOSED / "*" / "meta.json"))):
        meta = json.loads(Path(meta_p).read_text())
        if meta["exam"] != exam or meta["run_id"].startswith(("PILOT", "smoke")):
            continue
        if len(meta.get("students", [])) < 500:          # pilots and subsets
            continue
        wb = ROOT / meta["workbook"]                            # repo-relative since the fold-in
        s = stats(wb, exam)
        if not s:
            continue
        temp = (meta.get("sampling") or {}).get("temperature")
        rows.append({"run_id": meta["run_id"], "model": meta["model"], "strictness": meta["persona"],
                     "prompt": prompt_flags_from_config(meta["config"]),
                     "temperature": None if temp is None else float(temp), **s})
    return sorted(rows, key=lambda r: r["run_id"])


def prompt_baseline_rows() -> list[dict]:
    rows = []
    for p in sorted(glob.glob(str(PB / "PB-*__*.xlsx"))):
        name = Path(p).name
        rid = name.split("__")[0]
        m = re.search(r"_m-([^_]+)_mode-([a-z]+)_(sol\d_gd\d_rs\d_bd\d)_str-([a-z]+)_t(\d\d)_", name)
        model, mode, config, persona, t = m.groups()
        s = stats(_with_rebuilt_totals(Path(p)), "cv")
        if not s:
            continue
        rows.append({"run_id": rid, "model": model, "strictness": f"{persona} ({mode})",
                     "prompt": prompt_flags_from_config(config), "temperature": int(t) / 10, **s})
    return sorted(rows, key=lambda r: r["run_id"])


def tex_rows(rows: list[dict], tex_escape) -> list[str]:
    out = []
    for r in rows:
        t = "---" if r["temperature"] is None else ("0" if r["temperature"] == 0 else f"{r['temperature']:.1f}")
        out.append(" & ".join([
            tex_escape(r["run_id"]), rf"\texttt{{{tex_escape(r['model'])}}}", tex_escape(r["strictness"]),
            r["prompt"], f"${t}$" if t != "---" else t, f"${r['n']}$", f"${r['mae']:.2f}$",
            f"$[{r['lo']:.2f},\\ {r['hi']:.2f}]$", f"${r['bias']:+.2f}$", tex_escape(r["behaviour"])]) + r" \\")
    return out


if __name__ == "__main__":
    for title, rows in (("CV: closed", closed_rows("cv")), ("CV: prompt baselines", prompt_baseline_rows()),
                        ("ML: closed", closed_rows("ml"))):
        print(f"## {title} ({len(rows)} rows)")
        for r in rows:
            print(f"  {r['run_id']:<7}{r['model']:<20}{r['strictness']:<22}{r['prompt']:<5}"
                  f"t={r['temperature']!s:<5} n={r['n']:<5} MAE {r['mae']:.2f} [{r['lo']:.2f}, {r['hi']:.2f}] "
                  f"bias {r['bias']:+.2f}  {r['behaviour']}")
