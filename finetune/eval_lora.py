#!/usr/bin/env python3
"""Evaluate a marks-only fine-tuned grader on the held-out split.

Loads a base Qwen model (optionally + a trained LoRA adapter), predicts
``{score, bonus}`` for every (held-out student, question), sums to a total, and
writes a result workbook in the SAME shape the ablation grader produces so
``analysis/computer_vision_run_analysis.py`` picks it up automatically. Also prints MAE vs the
TA-average total over the eval students.

Pass ``--adapter ''`` (or omit it) to evaluate the BASE model -> that is your
apples-to-apples baseline on the identical eval split.

Run on a GPU node:
    python finetune/eval_lora.py \
        --model Qwen/Qwen2.5-Coder-7B-Instruct \
        --adapter finetune/adapters/qwen2.5-coder-7b \
        --eval-students finetune/data/eval_students.json \
        --tag FT-A07
"""
from __future__ import annotations

import argparse
import json
import shutil
import sys
import time
from pathlib import Path

import torch
import openpyxl
from transformers import AutoModelForCausalLM, AutoTokenizer, BitsAndBytesConfig

HERE = Path(__file__).resolve().parent
ROOT = HERE.parent
sys.path.insert(0, str(HERE))  # vendored common.py lives beside this script

from common import (  # noqa: E402
    QUESTION_META,
    SOLUTIONS_DIR,
    EXTRACTED_DIR,
    GRADES_XLSX,
    GUIDELINES,
    RESULTS_DIR,
    load_rubric_sections,
    notebook_to_text,
    ai_headers,
    ensure_ai_columns,
    build_student_row_index,
    atomic_save,
    compute_run_summary,
    _parse_json_lenient,
    build_prompt,
)
from build_finetune_data import build_scores_prompt, truncate_code  # noqa: E402
from build_breakdown_data import build_bd_prompt  # noqa: E402


def parse_args() -> argparse.Namespace:
    ap = argparse.ArgumentParser(description=__doc__,
                                 formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--model", required=True, help="Base HuggingFace model id.")
    ap.add_argument("--adapter", default="",
                    help="Path to the trained LoRA adapter. Empty = base model (baseline).")
    ap.add_argument("--eval-students", type=Path,
                    default=HERE / "data" / "eval_students.json")
    ap.add_argument("--tag", required=True,
                    help="Run tag -> finetune/results/<tag>.xlsx.")
    ap.add_argument("--output-xlsx", type=Path, default=None)
    ap.add_argument("--no-solution", action="store_true")
    ap.add_argument("--no-guidelines", action="store_true")
    ap.add_argument("--strictness", default="neutral")
    ap.add_argument("--max-student-chars", type=int, default=24000)
    ap.add_argument("--prompt-style",
                    choices=["marks", "breakdown", "breakdown_reasoning"], default="marks",
                    help="Must match how the adapter was trained. "
                         "marks               = {score,bonus} only (finetune/data). "
                         "breakdown           = {score,bonus,task_breakdown}, NO reasoning "
                         "(D01 distil, finetune/data_bd -- the default bd build). "
                         "breakdown_reasoning = full bd1 JSON incl. reasoning "
                         "(only if you built data_bd with --with-reasoning).")
    ap.add_argument("--questions", default="1,2,3,4",
                    help="Only grade these questions, e.g. '3' for a leave-Q3-out test. "
                         "Reports per-question MAE (AI Q vs TA Q) for each.")
    ap.add_argument("--max-new-tokens", type=int, default=0,
                    help="0 = auto (64 for marks, 1024 for breakdown).")
    ap.add_argument("--batch-size", type=int, default=1,
                    help="Prompts per generate() call. >1 batches the HF eval so the GPU "
                         "isn't left ~80%% idle at batch size 1 (roughly Bx faster). Uses "
                         "left-padding; results match batch-1 up to the usual "
                         "floating-point/batching noise. B=8 is safe on an 80 GB card for "
                         "an 8-14B model; lower it if you OOM.")
    ap.add_argument("--no-4bit", action="store_true")
    ap.add_argument("--attn", default="sdpa", choices=["sdpa", "eager", "flash_attention_2"])
    return ap.parse_args()


def load_model(model_id: str, adapter: str, four_bit: bool, attn: str = "sdpa"):
    tok = AutoTokenizer.from_pretrained(model_id, trust_remote_code=True)
    if tok.pad_token is None:
        tok.pad_token = tok.eos_token
    quant = None
    if four_bit:
        quant = BitsAndBytesConfig(
            load_in_4bit=True,
            bnb_4bit_quant_type="nf4",
            bnb_4bit_use_double_quant=True,
            bnb_4bit_compute_dtype=torch.bfloat16,
        )
    model = AutoModelForCausalLM.from_pretrained(
        model_id,
        quantization_config=quant,
        dtype=torch.bfloat16,
        device_map={"": 0},
        trust_remote_code=True,
        attn_implementation=attn,
    )
    if adapter:
        from peft import PeftModel
        model = PeftModel.from_pretrained(model, adapter)
        print(f"  loaded adapter: {adapter}")
    model.eval()
    return model, tok


def _parse_or_recover(gen: str) -> dict:
    try:
        return _parse_json_lenient(gen)
    except Exception:
        s, e = gen.find("{"), gen.rfind("}")   # slice out first {...} (fenced/prose)
        if 0 <= s < e:
            try:
                return _parse_json_lenient(gen[s:e + 1])
            except Exception:
                pass
        return {}


def _render(tok, prompt: str) -> str:
    """Chat-template a single user prompt with thinking DISABLED, so the model
    emits the JSON directly instead of a reasoning preamble that can overrun the
    token budget. Templates that don't accept the kwarg raise TypeError -> fall
    back to the plain call."""
    msgs = [{"role": "user", "content": prompt}]
    try:
        return tok.apply_chat_template(msgs, tokenize=False,
                                       add_generation_prompt=True,
                                       enable_thinking=False)
    except TypeError:
        return tok.apply_chat_template(msgs, tokenize=False,
                                       add_generation_prompt=True)


@torch.no_grad()
def predict(model, tok, prompt: str, max_new_tokens: int) -> dict:
    text = _render(tok, prompt)
    inputs = tok(text, return_tensors="pt").to(model.device)
    out = model.generate(
        **inputs,
        max_new_tokens=max_new_tokens,
        do_sample=False,
        temperature=None,
        top_p=None,
        pad_token_id=tok.pad_token_id,
    )
    gen = tok.decode(out[0][inputs["input_ids"].shape[1]:], skip_special_tokens=True)
    return _parse_or_recover(gen)


@torch.no_grad()
def predict_batch(model, tok, prompts: list[str], max_new_tokens: int) -> list[dict]:
    """Greedy-generate for a batch of prompts at once. Left-padding is required for
    decoder-only generation so every sequence's real tokens end at the same
    position; each row's completion is then the tokens after the (shared) padded
    input length. Equivalent to calling predict() per prompt, up to floating-point
    batching noise."""
    if len(prompts) == 1:
        return [predict(model, tok, prompts[0], max_new_tokens)]
    texts = [_render(tok, p) for p in prompts]
    old_side = tok.padding_side
    tok.padding_side = "left"
    try:
        enc = tok(texts, return_tensors="pt", padding=True).to(model.device)
    finally:
        tok.padding_side = old_side
    out = model.generate(
        **enc,
        max_new_tokens=max_new_tokens,
        do_sample=False,
        temperature=None,
        top_p=None,
        pad_token_id=tok.pad_token_id,
    )
    in_len = enc["input_ids"].shape[1]          # same for all rows (left-padded)
    gens = tok.batch_decode(out[:, in_len:], skip_special_tokens=True)
    return [_parse_or_recover(g) for g in gens]


def main() -> None:
    args = parse_args()
    if not torch.cuda.is_available():
        raise SystemExit("No CUDA device visible. Run on a GPU node.")

    # Breakdown-only completions measure up to ~3.1k chars (~890 tokens) in the
    # distil data, so 1024 leaves headroom without paying the reasoning tax.
    auto_tokens = {"marks": 64, "breakdown": 1024, "breakdown_reasoning": 1536}
    max_new_tokens = args.max_new_tokens or auto_tokens[args.prompt_style]
    questions = sorted(int(x) for x in args.questions.split(",") if x.strip())
    print(f"Prompt style: {args.prompt_style}  (max_new_tokens={max_new_tokens})  "
          f"questions={questions}")

    eval_students = json.loads(Path(args.eval_students).read_text())
    print(f"Eval students: {len(eval_students)}")

    rubric_sections = load_rubric_sections()
    guidelines = GUIDELINES.read_text(encoding="utf-8")
    use_solution = not args.no_solution
    use_guidelines = not args.no_guidelines
    solutions = {}
    for q in (1, 2, 3, 4):
        sp = SOLUTIONS_DIR / QUESTION_META[q][2]
        solutions[q] = notebook_to_text(sp) if (use_solution and sp.exists()) else None

    model, tok = load_model(args.model, args.adapter, not args.no_4bit, args.attn)

    out_xlsx = (args.output_xlsx.resolve() if args.output_xlsx
                else (RESULTS_DIR / f"{args.tag}.xlsx").resolve())
    out_xlsx.parent.mkdir(parents=True, exist_ok=True)
    shutil.copy2(GRADES_XLSX, out_xlsx)
    wb = openpyxl.load_workbook(out_xlsx)
    ws = wb.active
    # Clear any pre-existing AI columns, then ensure our headers exist.
    hdrs = [c.value for c in ws[1]]
    for r in range(2, ws.max_row + 1):
        for i, h in enumerate(hdrs):
            if h and str(h).startswith("AI "):
                ws.cell(row=r, column=i + 1).value = None
    cols = ensure_ai_columns(ws)
    row_index = build_student_row_index(ws)
    atomic_save(wb, out_xlsx)

    def make_prompt(q: int, code: str) -> str:
        if args.prompt_style == "breakdown":            # breakdown-only, no reasoning
            return build_bd_prompt(q, rubric_sections[q], guidelines, code,
                                   solutions[q], use_guidelines, args.strictness)
        if args.prompt_style == "breakdown_reasoning":
            return build_prompt(q=q, rubric_section=rubric_sections[q],
                                guidelines=guidelines, student_code=code,
                                solution_code=solutions[q], use_breakdown=True,
                                use_guidelines=use_guidelines, strictness=args.strictness)
        return build_scores_prompt(q, rubric_sections[q], guidelines, code,
                                   solutions[q], use_guidelines, args.strictness)

    # --- Phase 1: build the work-list; record no-submission cells up front ---
    # per_student[student][q] = (score, bonus, reason). Missing notebooks are
    # filled here as 0; the rest become generation jobs.
    per_student: dict[int, dict[int, tuple]] = {}
    jobs = []  # (student, q, prompt)
    for student in eval_students:
        if row_index.get(student) is None:
            continue
        per_student[student] = {}
        sdir = EXTRACTED_DIR / str(student)
        for q in questions:
            nb = sdir / f"Q{q}.ipynb"
            if not nb.exists():
                per_student[student][q] = (0.0, 0.0, "No submission found.")
                continue
            code = truncate_code(notebook_to_text(nb), args.max_student_chars)
            jobs.append((student, q, make_prompt(q, code)))
    print(f"  {len(jobs)} generations to run (batch size {args.batch_size})", flush=True)

    # --- Phase 2: batch-generate; write cells as results arrive ---
    t0 = time.time()
    done = 0
    for b in range(0, len(jobs), args.batch_size):
        chunk = jobs[b:b + args.batch_size]
        results = predict_batch(model, tok, [p for _, _, p in chunk], max_new_tokens)
        for (student, q, _), res in zip(chunk, results):
            max_score, max_bonus, _ = QUESTION_META[q]
            score = max(0.0, min(float(res.get("score", 0) or 0), max_score))
            bonus = max(0.0, min(float(res.get("bonus", 0) or 0), max_bonus))
            per_student[student][q] = (score, bonus, f"(FT {args.prompt_style})")
            row = row_index[student]
            ws.cell(row=row, column=cols[f"AI Q{q} Score"], value=score)
            bcol = cols.get(f"AI Q{q} Bonus")
            if bcol:
                ws.cell(row=row, column=bcol, value=bonus)
            ws.cell(row=row, column=cols[f"AI Q{q} Reasoning"], value=f"(FT {args.prompt_style})")
        done += len(chunk)
        rate = done / (time.time() - t0 + 1e-9)
        eta = (len(jobs) - done) / rate / 60 if rate else 0
        print(f"  {done}/{len(jobs)} gens  {rate:.2f}/s  ETA {eta:.1f} min", flush=True)
        if (b // max(1, args.batch_size)) % 10 == 0:
            atomic_save(wb, out_xlsx)

    # --- Phase 3: write per-student totals (incl. no-submission zeros) ---
    for student, qmap in per_student.items():
        row = row_index[student]
        # ensure no-submission cells are written too
        for q, (sc, bo, rz) in qmap.items():
            ws.cell(row=row, column=cols[f"AI Q{q} Score"], value=sc)
            bcol = cols.get(f"AI Q{q} Bonus")
            if bcol:
                ws.cell(row=row, column=bcol, value=bo)
            ws.cell(row=row, column=cols[f"AI Q{q} Reasoning"], value=rz)
        tot_s = sum(v[0] for v in qmap.values())
        tot_b = sum(v[1] for v in qmap.values())
        ws.cell(row=row, column=cols["AI Total Score"], value=tot_s)
        ws.cell(row=row, column=cols["AI Total Bonus"], value=tot_b)
    atomic_save(wb, out_xlsx)

    summary = compute_run_summary(out_xlsx)
    print("\n== Eval summary (held-out split) ==")
    print(f"  adapter      : {args.adapter or '(base model, no adapter)'}")
    print(f"  n_graded     : {summary['n_graded']}")
    print(f"  mean total   : {summary['mean_total_score']}")
    print(f"  MAE vs TA-avg: {summary['MAE_vs_TA_avg']}")
    print(f"  output       : {out_xlsx}")

    # Per-question MAE: AI Q{q} Score vs the TA-average of that question's score.
    # This is the metric that matters for a leave-one-question-out test.
    wbq = openpyxl.load_workbook(out_xlsx, data_only=True)
    wsq = wbq.active
    Hq = [c.value for c in wsq[1]]
    def qcol(name):
        return Hq.index(name) + 1 if name in Hq else None
    print("  per-question MAE (AI Qk vs TA Qk avg):")
    for q in questions:
        c_ai = qcol(f"AI Q{q} Score")
        c_ta = [qcol(f"TA {t} - Q{q} Score") for t in (1, 2)]
        if c_ai is None:
            continue
        diffs = []
        for r in wsq.iter_rows(min_row=2, values_only=True):
            v = r[c_ai - 1]
            if v in (None, ""):
                continue
            tv = [r[c - 1] for c in c_ta if c and isinstance(r[c - 1], (int, float))]
            if tv:
                diffs.append(float(v) - sum(tv) / len(tv))
        if diffs:
            mae = sum(abs(d) for d in diffs) / len(diffs)
            bias = sum(diffs) / len(diffs)
            print(f"    Q{q}: n={len(diffs)}  MAE={mae:.3f}  bias={bias:+.3f}")


if __name__ == "__main__":
    main()
