"""Exam registry: the per-course facts the finetune pipeline varies over.

Two exams live in this project and they differ in almost every surface detail --
question count, score caps, rubric header style, ground-truth file format, and
whether an exam-guidelines document exists at all:

                     advance (cv)                 intro
    students         570                          1038
    questions        Q1..Q4                       Q1..Q3
    base + bonus     35.0 + 13.0                  56.0 + 9.0
    ground truth     .xlsx, Score/Bonus split     .csv, one folded Grade
    rubric headers   '## Question 1: ...'         '## Q1: ...'
    guidelines       exam_guidelines_student.md   (none)

Everything that differs is captured in an ExamSpec so the builders, trainer and
eval scripts stay exam-agnostic and take a --dataset flag instead.

The 'both' variant is NOT an ExamSpec -- it is a pooled *split* over the two
real exams (see build_finetune_data.make_pooled_split). Pooling happens at the
student-list level so each exam keeps its own prompts, caps and rubric.
"""
from __future__ import annotations

import csv
import json
import re
from dataclasses import dataclass
from pathlib import Path
from typing import Optional

HERE = Path(__file__).resolve().parent
ROOT = HERE.parent


@dataclass(frozen=True)
class ExamSpec:
    key: str
    label: str                 # goes into the prompt's first line
    dataset_dir: Path
    rubric_path: Path
    rubric_header_re: str      # splits the rubric into per-question sections
    solutions_dir: Path
    extracted_dir: Path
    grades_path: Path
    grades_format: str         # "xlsx" | "csv"
    ta_mode: str               # "split" -> Score/Bonus cols | "combined" -> one Grade col
    question_meta: dict[int, tuple[float, float, str]]
    teacher_xlsx: Path         # bd1 teacher run distilled for breakdown training
    guidelines_path: Optional[Path] = None
    # intro's rubric ends each question with a '---' page rule that is furniture,
    # not rubric. cv's does not, and must NOT be stripped: doing so would shift
    # its prompts off the frozen ones behind the published FT-* runs.
    strip_trailing_rule: bool = False

    @property
    def questions(self) -> tuple[int, ...]:
        return tuple(sorted(self.question_meta))

    @property
    def total_base(self) -> float:
        return sum(m[0] for m in self.question_meta.values())

    @property
    def total_bonus(self) -> float:
        return sum(m[1] for m in self.question_meta.values())

    @property
    def has_guidelines(self) -> bool:
        return self.guidelines_path is not None and self.guidelines_path.exists()


CV = ExamSpec(
    key="cv",
    label="KAUST Stage 3 Practical AI",
    dataset_dir=ROOT / "computer_vision_dataset",
    rubric_path=ROOT / "computer_vision_dataset" / "stage_3_rubrics_ta.md",
    # '## Question 1: ...' and '## Bonus Question 4: ...'
    rubric_header_re=r"(?m)^(##\s+.*)$",
    solutions_dir=ROOT / "computer_vision_dataset" / "Solutions",
    extracted_dir=ROOT / "computer_vision_dataset" / "submissions_extracted",
    grades_path=ROOT / "computer_vision_dataset" / "Practical_AI_exam_grades.xlsx",
    grades_format="xlsx",
    ta_mode="split",
    question_meta={
        1: (12.0, 4.0, "Q1_stage3_Finetuning_Solution.ipynb"),
        2: (11.0, 4.0, "Q2_stage3_CNN_Solution.ipynb"),
        3: (12.0, 0.0, "Q3_stage3_Segmentation_Solution.ipynb"),
        4: (0.0, 5.0, "Q4_stage3_Colorization_Solution.ipynb"),
    },
    teacher_xlsx=(
        ROOT / "computer_vision_results" / "results"
        / "D01__D01_m-3-flash-preview_sol1_gd1_rs0_bd1_str-neutral_t00_n1.xlsx"
    ),
    guidelines_path=ROOT / "computer_vision_dataset" / "exam_guidelines_student.md",
)

INTRO = ExamSpec(
    key="intro",
    label="KAUST Stage 2 Introduction to AI Practical",
    dataset_dir=ROOT / "introduction_to_ai_dataset",
    rubric_path=ROOT / "introduction_to_ai_dataset" / "stage_2_rubrics_ta.md",
    # '## Q1: Regression -- ...'; the plain '##' split would also grab the title.
    rubric_header_re=r"(?m)^(##\s+Q\d.*)$",
    solutions_dir=ROOT / "introduction_to_ai_dataset" / "Solutions",
    extracted_dir=ROOT / "introduction_to_ai_dataset" / "submissions_extracted",
    grades_path=ROOT / "introduction_to_ai_dataset" / "Practical_AI_exam_grades.csv",
    grades_format="csv",
    ta_mode="combined",
    # Caps from the marking scheme's task tables, matching
    # introduction_to_ai_grade_with_local.py QUESTION_META exactly.
    question_meta={
        1: (23.0, 3.0, "Q1_Regression_Solution.ipynb"),
        2: (14.0, 3.0, "Q2_Pytorch_Solution.ipynb"),
        3: (19.0, 3.0, "Q3_Classification_Solution.ipynb"),
    },
    teacher_xlsx=(
        ROOT / "introduction_to_ai_results" / "results"
        / "IG07__IG07_m-3-flash-preview_sol1_gd0_rs0_bd1_str-neutral_t00_n1.xlsx"
    ),
    guidelines_path=None,      # this exam shipped no student-guidelines document
    strip_trailing_rule=True,
)

EXAMS: dict[str, ExamSpec] = {CV.key: CV, INTRO.key: INTRO}

# --dataset accepts these. "both" is handled by the callers as a pooled split
# over the two real exams, not as an ExamSpec of its own.
DATASET_CHOICES = ("cv", "intro", "both")


def get_exam(key: str) -> ExamSpec:
    try:
        return EXAMS[key]
    except KeyError:
        raise SystemExit(
            f"unknown exam {key!r}; expected one of {sorted(EXAMS)}"
        ) from None


def exams_for(dataset: str) -> list[ExamSpec]:
    """--dataset value -> the exams it spans. 'both' -> [CV, INTRO]."""
    if dataset == "both":
        return [CV, INTRO]
    return [get_exam(dataset)]


# ---------------------------------------------------------------------------
# Rubric / guidelines
# ---------------------------------------------------------------------------

def load_rubric_sections(exam: ExamSpec) -> dict[int, str]:
    """q -> that question's rubric section, split on the exam's header style."""
    text = exam.rubric_path.read_text(encoding="utf-8")
    parts = re.split(exam.rubric_header_re, text)
    sections: dict[int, str] = {}
    for i in range(1, len(parts), 2):
        header = parts[i].strip()
        body = parts[i + 1] if i + 1 < len(parts) else ""
        m = re.match(r"##\s+(?:Bonus\s+)?(?:Question\s+|Q)(\d+)", header, re.IGNORECASE)
        if not m:
            continue
        sec = header + "\n" + body.strip()
        if exam.strip_trailing_rule:
            sec = sec.rstrip("- \n")
        sections[int(m.group(1))] = sec
    missing = [q for q in exam.questions if q not in sections]
    if missing:
        raise RuntimeError(f"{exam.key}: rubric sections missing for {missing}")
    return sections


def load_guidelines(exam: ExamSpec) -> str:
    """The exam-guidelines text, or '' when the exam has none (intro)."""
    if not exam.has_guidelines:
        return ""
    return exam.guidelines_path.read_text(encoding="utf-8")


# ---------------------------------------------------------------------------
# Ground truth
# ---------------------------------------------------------------------------

def _num(v) -> Optional[float]:
    if isinstance(v, (int, float)):
        return float(v)
    if isinstance(v, str):
        v = v.strip()
        if v:
            try:
                return float(v)
            except ValueError:
                return None
    return None


def _split_combined(grade: float, max_score: float, max_bonus: float) -> tuple[float, float]:
    """Intro folds bonus into one per-question Grade; recover a (score, bonus) pair.

    Base is filled first and the overflow becomes bonus, so score + bonus == grade
    whenever grade <= max_score + max_bonus. That matters because
    introduction_to_ai_grade_with_local.py compares (AI score + AI bonus) against
    the TA total -- the metric only sees the sum, so how we attribute within it
    cannot bias the result, while keeping both caps the prompt promises.

    Marks above base+bonus (~0.5% of intro TA marks) are clamped, exactly as the
    root grader clamps model output to the rubric.
    """
    score = max(0.0, min(grade, max_score))
    bonus = max(0.0, min(grade - score, max_bonus))
    return score, bonus


def _ta_rows(exam: ExamSpec):
    """Yield ground-truth rows as {header: value} dicts, xlsx or csv alike."""
    if exam.grades_format == "csv":
        with open(exam.grades_path, newline="", encoding="utf-8-sig") as f:
            yield from csv.DictReader(f)
        return
    import openpyxl
    ws = openpyxl.load_workbook(exam.grades_path, data_only=True).active
    headers = [c.value for c in ws[1]]
    for row in ws.iter_rows(min_row=2, values_only=True):
        yield {h: v for h, v in zip(headers, row) if h is not None}


def load_ta_targets(exam: ExamSpec) -> dict[int, dict[int, dict[str, float]]]:
    """student -> q -> {'score', 'bonus'}, TA1/TA2 averaged and clamped to the caps.

    A question is skipped only when no TA supplied a usable mark for a component
    the question actually has (cv Q4 has no base, Q3 no bonus -> those read 0.0).
    """
    out: dict[int, dict[int, dict[str, float]]] = {}
    for row in _ta_rows(exam):
        num = _num(row.get("Number"))
        if num is None:
            continue
        student = int(num)
        per_q: dict[int, dict[str, float]] = {}
        for q in exam.questions:
            max_score, max_bonus, _ = exam.question_meta[q]

            def avg(field: str) -> Optional[float]:
                vals = [
                    v for ta in (1, 2)
                    if (v := _num(row.get(f"TA {ta} - Q{q} {field}"))) is not None
                ]
                return sum(vals) / len(vals) if vals else None

            if exam.ta_mode == "combined":
                grade = avg("Grade")
                if grade is None:
                    continue
                score, bonus = _split_combined(grade, max_score, max_bonus)
            else:
                score = avg("Score") if max_score > 0 else 0.0
                bonus = avg("Bonus") if max_bonus > 0 else 0.0
                if max_score > 0 and score is None:
                    continue
                score = max(0.0, min(score if score is not None else 0.0, max_score))
                bonus = max(0.0, min(bonus if bonus is not None else 0.0, max_bonus))
            per_q[q] = {"score": round(score, 2), "bonus": round(bonus, 2)}
        if per_q:
            out[student] = per_q
    return out


def load_solutions(exam: ExamSpec, use_solution: bool, to_text) -> dict[int, Optional[str]]:
    """q -> reference-solution text (or None when --no-solution / file absent)."""
    if not use_solution:
        return {q: None for q in exam.questions}
    sols: dict[int, Optional[str]] = {}
    for q in exam.questions:
        p = exam.solutions_dir / exam.question_meta[q][2]
        sols[q] = to_text(p) if p.exists() else None
    return sols


def ai_headers(exam: ExamSpec) -> list[str]:
    cols: list[str] = []
    for q in exam.questions:
        cols.append(f"AI Q{q} Score")
        if exam.question_meta[q][1] > 0:
            cols.append(f"AI Q{q} Bonus")
        cols.append(f"AI Q{q} Reasoning")
    return cols + ["AI Total Score", "AI Total Bonus"]


def describe(exam: ExamSpec) -> str:
    return (
        f"{exam.key}: Q{list(exam.questions)} "
        f"base={exam.total_base} bonus={exam.total_bonus} "
        f"gt={exam.grades_path.name} ({exam.ta_mode}) "
        f"guidelines={'yes' if exam.has_guidelines else 'no'}"
    )


# ---------------------------------------------------------------------------
# Result workbooks
# ---------------------------------------------------------------------------

def _coerce_num(v: str):
    if v == "":
        return None
    try:
        f = float(v)
        return int(f) if f.is_integer() and "." not in v else f
    except (ValueError, TypeError):
        return v


def init_results_workbook(exam: ExamSpec, dest: Path) -> None:
    """Seed a results workbook with the exam's ground truth.

    cv has a master .xlsx (with formula columns) that is simply copied, mirroring
    what eval_vllm did before. intro released plain-value CSV instead, so the
    workbook is built from its rows -- the same approach as
    introduction_to_ai_grade_with_local.init_output_workbook. AI columns are
    appended afterwards by common.ensure_ai_columns.
    """
    import shutil
    dest.parent.mkdir(parents=True, exist_ok=True)
    if exam.grades_format == "xlsx":
        shutil.copy2(exam.grades_path, dest)
        return
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    with open(exam.grades_path, newline="", encoding="utf-8-sig") as f:
        for row in csv.reader(f):
            ws.append([_coerce_num(v) for v in row])
    tmp = dest.with_name(dest.name + ".tmp")
    wb.save(tmp)
    import os
    os.replace(tmp, dest)


def ta_total_columns(exam: ExamSpec) -> dict[int, dict]:
    """Per-TA column names needed to rebuild a TA total from the workbook.

    cv: 'TA n - Total Score (out of 35)' with per-question Score columns as a
        fallback when the total cell is an uncached formula.
    intro: no total-score column of that name and bonus is folded into each
        per-question Grade, so the total is the sum of the Grade columns and the
        AI side must be compared as (score + bonus). See
        introduction_to_ai_grade_with_local.py's module docstring.
    """
    if exam.ta_mode == "combined":
        return {ta: {"total": f"TA {ta} - Total Grade",
                     "per_q": [f"TA {ta} - Q{q} Grade" for q in exam.questions],
                     "ai_includes_bonus": True}
                for ta in (1, 2)}
    return {ta: {"total": f"TA {ta} - Total Score (out of {exam.total_base:g})",
                 "per_q": [f"TA {ta} - Q{q} Score" for q in exam.questions],
                 "ai_includes_bonus": False}
            for ta in (1, 2)}


def run_summary(exam: ExamSpec, output_xlsx: Path) -> dict:
    """mean/std of the AI total and MAE against the TA average, exam-aware.

    Generalizes common.compute_run_summary, which hardcodes cv's four questions
    and its 'out of 35' header. For intro the AI total compared against the TAs
    is score + bonus, because intro's ground truth folds bonus into the grade.
    """
    import logging
    import statistics
    import openpyxl
    out = {"mean_total_score": "", "std_total_score": "",
           "MAE_vs_TA_avg": "", "n_graded": 0}
    if not output_xlsx.exists():
        return out
    try:
        ws = openpyxl.load_workbook(output_xlsx, data_only=True).active
        H = [c.value for c in ws[1]]
        idx = {h: i for i, h in enumerate(H) if h is not None}
        c_ai = idx.get("AI Total Score")
        c_aib = idx.get("AI Total Bonus")
        if c_ai is None:
            return out
        spec = ta_total_columns(exam)
        include_bonus = spec[1]["ai_includes_bonus"]

        def ta_total(row, ta):
            s = spec[ta]
            v = row[idx[s["total"]]] if s["total"] in idx else None
            if isinstance(v, (int, float)):
                return float(v)
            parts = [row[idx[c]] for c in s["per_q"] if c in idx]
            parts = [p for p in parts if isinstance(p, (int, float))]
            return float(sum(parts)) if parts else None

        scores, diffs = [], []
        for row in ws.iter_rows(min_row=2, values_only=True):
            v = row[c_ai]
            if v in (None, ""):
                continue
            try:
                ai = float(v)
            except (TypeError, ValueError):
                continue
            if include_bonus and c_aib is not None:
                b = row[c_aib]
                if isinstance(b, (int, float)):
                    ai += float(b)
            scores.append(ai)
            tas = [t for t in (ta_total(row, 1), ta_total(row, 2)) if t is not None]
            if tas:
                diffs.append(ai - sum(tas) / len(tas))
        out["n_graded"] = len(scores)
        if scores:
            out["mean_total_score"] = round(statistics.fmean(scores), 3)
            if len(scores) >= 2:
                out["std_total_score"] = round(statistics.stdev(scores), 3)
        if diffs:
            out["MAE_vs_TA_avg"] = round(sum(abs(d) for d in diffs) / len(diffs), 3)
    except Exception as e:  # noqa: BLE001
        logging.warning("Failed to summarize %s: %s", output_xlsx, e)
    return out


def load_eval_students(path: Path, exam: ExamSpec) -> list[int]:
    """Read eval_students.json, which is a flat list for a single-exam build
    and {exam_key: [...]} for a pooled ('both') one."""
    data = json.loads(Path(path).read_text())
    if isinstance(data, dict):
        if exam.key not in data:
            raise SystemExit(
                f"{path} has no '{exam.key}' entry (has {sorted(data)})")
        return data[exam.key]
    return data
