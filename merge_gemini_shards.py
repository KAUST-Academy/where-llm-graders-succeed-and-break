#!/usr/bin/env python3
"""merge_gemini_shards.py -- fold sharded Gemini-grader workbooks into one run workbook.

The Gemini graders are sequential (one call in flight), so a full-cohort run
is split into `--students a-b` shards, each writing its own clone of the grades
template. Every clone has identical rows, so merging is row-wise: for each
student row that a shard graded (any `AI Q* Score` set), copy that shard's AI
columns into the target workbook. The target may already hold rows graded by
an earlier sequential run; shards only fill rows whose AI cells are still empty
unless --overwrite is given.

    python merge_gemini_shards.py --target computer_vision_results/results/F03__...xlsx \
        --shards closed_runs/F03_shards/shard*.xlsx
"""
import argparse
import glob
import shutil
from pathlib import Path

import openpyxl


def ai_columns(ws) -> dict[str, int]:
    return {c.value: i + 1 for i, c in enumerate(ws[1]) if c.value and str(c.value).startswith("AI ")}


def student_rows(ws) -> dict[int, int]:
    out = {}
    for row in ws.iter_rows(min_row=2):
        v = row[0].value
        try:
            out[int(v)] = row[0].row
        except (TypeError, ValueError):
            continue
    return out


def main() -> None:
    ap = argparse.ArgumentParser()
    ap.add_argument("--target", required=True)
    ap.add_argument("--shards", nargs="+", required=True)
    ap.add_argument("--overwrite", action="store_true")
    a = ap.parse_args()
    target = Path(a.target)
    shards = sorted(p for pat in a.shards for p in glob.glob(pat))
    if not shards:
        raise SystemExit("no shard files")
    if not target.exists():                       # start from the first shard's clone
        shutil.copy2(shards[0], target)
    wb = openpyxl.load_workbook(target)
    ws = wb.active
    tcols = ai_columns(ws)
    trows = student_rows(ws)
    score_cols = [c for h, c in tcols.items() if h.endswith(" Score") and h != "AI Total Score"]
    merged = skipped = 0
    for sp in shards:
        sws = openpyxl.load_workbook(sp, data_only=True).active
        scols = ai_columns(sws)
        srows = student_rows(sws)
        if set(scols) != set(tcols):
            raise SystemExit(f"{sp}: AI columns differ from target")
        for s, srow in srows.items():
            if not any(sws.cell(row=srow, column=scols[h]).value not in (None, "")
                       for h in tcols if h.endswith(" Score") and h != "AI Total Score"):
                continue                          # shard did not grade this student
            trow = trows.get(s)
            if trow is None:
                continue
            already = any(ws.cell(row=trow, column=c).value not in (None, "") for c in score_cols)
            if already and not a.overwrite:
                skipped += 1
                continue
            for h, c in tcols.items():
                ws.cell(row=trow, column=c).value = sws.cell(row=srow, column=scols[h]).value
            merged += 1
    tmp = target.with_name(target.name + ".tmp")
    wb.save(tmp)
    tmp.replace(target)
    n_total = sum(1 for r in trows.values() if ws.cell(row=r, column=tcols["AI Total Score"]).value not in (None, ""))
    print(f"merged {merged} students from {len(shards)} shards ({skipped} already graded, kept); "
          f"target now has {n_total}/{len(trows)} AI totals -> {target}")


if __name__ == "__main__":
    main()
