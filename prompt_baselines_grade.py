#!/usr/bin/env python3
"""prompt_baselines_grade.py -- prompt-only baselines for the
strict-persona collapse (no weight updates), on the CV exam.

The reviewer asks whether prompt structuring alone repairs the collapse before
concluding that LoRA is needed. Two non-parametric strategies are implemented,
both under the paper's *strict* persona and on the paper's own prompt scaffold
(``computer_vision_grade_with_local.build_prompt`` and its text blocks are
reused verbatim; the paper's grader is not modified):

  decomposed   one call per rubric task row ("atomic per-criterion passes",
               cf. Rubric Is All You Need).  Each call carries the persona,
               the exam guidelines, the reference solution and the student's
               notebook exactly as the paper's prompt does, but the RUBRIC
               section holds a single task row (with its part heading and
               point value) and the model is asked for that row's award only:
               {"awarded": number, "reasoning": str}.  The question's base
               score is the sum of the awarded values over its base and
               penalty rows, clamped to [0, max base]; bonus rows sum to the
               bonus, clamped to [0, max bonus].  Penalty rows return the
               penalty (e.g. -2.0) when it applies and 0 otherwise.

  arbitrated   grade-then-arbitrate ("multi-agent deliberation", cf. ChatEval).
               Per question: (A) the paper's exact strict prompt, (B) the
               paper's exact neutral prompt, then (C) an arbiter call under the
               neutral persona that receives the rubric, both JSON grades, the
               reference solution and the submission, and must resolve every
               disagreement by the rubric's partial-credit scale, returning the
               paper's JSON contract (score, bonus, task_breakdown, reasoning).
               The arbiter's grade is the run's grade; A and B are kept.

Students default to 1-100 (the G-series subset), questions to Q1-Q3 (the
35-point base scale every metric in the paper uses; Q4 is bonus-only).

Backends
  --api-base URL        a vLLM OpenAI-compatible server (the paper's stack);
                        sampling mirrors the paper's grader: temperature 0,
                        top_p 1, top_k -1, thinking off; no guided_json (the paper's
                        vLLM 0.19.1 stack ignored it, Appendix A).
  --provider openai     the OpenAI API through closed_batch_grade's client
                        (key GPT= in .env), reasoning_effort none, temperature
                        0 -- for smoke-testing the prompt plumbing for cents
                        before any GPU is queued.

Outputs (all under prompt_baselines/, outside the analysis pipeline's trees)
  results/<RID>__<tag>.xlsx   the run workbook in the standard ablation shape
                              (cloned from the grades workbook; AI Q1..Q4
                              Score/Bonus/Reasoning, AI Total Score/Bonus);
                              resumable -- graded cells are skipped
  calls/<RID>.jsonl           one record per model call: student, question,
                              stage/row, prompt sha1 + length, raw reply,
                              parsed value, finish reason, latency
  Use `report` to compare a finished run with the paper's runs on the same
  students (MAE, bias, zero-total rate, behaviour class, per-question MAE).

Examples
  # smoke (OpenAI, one student, one question, first 3 rubric rows)
  python prompt_baselines_grade.py grade --mode decomposed --provider openai --model gpt-5.5 \
      --students 1 --questions 1 --max-rows 3 --run-id SMOKE-D
  python prompt_baselines_grade.py grade --mode arbitrated --provider openai --model gpt-5.5 \
      --students 1 --questions 1 --run-id SMOKE-A
  # real run against a vLLM server
  python prompt_baselines_grade.py grade --mode decomposed --api-base http://localhost:8000/v1 \
      --model Qwen/Qwen2.5-Coder-32B-Instruct --model-tag qwen2.5-coder-32b --run-id PB-D01
  python prompt_baselines_grade.py report --run-id PB-D01 --paper-strict L-C01 --paper-neutral L-A32
"""
from __future__ import annotations

import argparse
import concurrent.futures as cf
import hashlib
import json
import logging
import re
import statistics
import sys
import threading
import time
from pathlib import Path
from typing import Any, Optional

import openpyxl

ROOT = Path(__file__).resolve().parent
sys.path.insert(0, str(ROOT))
import computer_vision_grade_with_local as G  # noqa: E402  (read-only reuse)

OUT_ROOT = ROOT / "prompt_baselines"
RESULTS = OUT_ROOT / "results"
CALLS = OUT_ROOT / "calls"
LOCK = threading.Lock()

# ---------------------------------------------------------------------------
# rubric decomposition
# ---------------------------------------------------------------------------

_POINTS_RE = re.compile(r"^\**\s*([+−–-]?)\s*(\d+(?:\.\d+)?)\s*\**$")


def parse_rubric_rows(section: str, q: int) -> list[dict]:
    """Split one question's rubric section into task rows.

    Returns dicts with: part (heading text), task, points (signed float),
    kind in {"base", "penalty", "bonus"}.  Total rows are dropped.
    """
    rows: list[dict] = []
    part = ""
    in_bonus_part = False
    for line in section.splitlines():
        h = re.match(r"^###\s+(.*)$", line)
        if h:
            part = h.group(1).strip()
            in_bonus_part = "bonus" in part.lower()
            continue
        if not line.startswith("|"):
            continue
        cells = [c.strip() for c in line.strip().strip("|").split("|")]
        if len(cells) < 2:
            continue
        task, pts = cells[0], cells[1]
        if re.match(r"^[-:\s]+$", task) or task.lower() in ("task", "bonus task"):
            continue                       # separator / header rows
        if "total" in task.lower() and task.startswith("**"):
            continue                       # "**Part 1 Total**" rows
        m = _POINTS_RE.match(pts)
        if not m:
            continue
        sign, num = m.group(1), float(m.group(2))
        points = -num if sign in ("−", "-", "–") else num
        if points < 0 or task.lower().startswith("penalty"):
            kind = "penalty"
        elif in_bonus_part or q == 4:
            kind = "bonus"
        else:
            kind = "base"
        rows.append({"part": part, "task": task, "points": points, "kind": kind})
    return rows


def rubric_item_block(q: int, row: dict, idx: int, n_rows: int) -> str:
    unit = "Bonus Task" if row["kind"] == "bonus" else "Task"
    pts = f"{row['points']:+g}" if row["kind"] == "penalty" else f"{row['points']:g}"
    return "\n".join([
        f"## RUBRIC ITEM {idx} OF {n_rows} FOR QUESTION {q}",
        f"Section: {row['part']}",
        "",
        f"| {unit} | Points |",
        "|------|-------:|",
        f"| {row['task']} | {pts} |",
        "",
        "(This is one item of the question's rubric. Decide this item only; the "
        "other items are graded in separate passes.)",
    ])


def build_decomposed_prompt(q: int, row: dict, idx: int, n_rows: int, guidelines: str,
                            student_code: str, solution_code: Optional[str],
                            strictness: str) -> str:
    """The paper's prompt scaffold with a one-row rubric and a one-value task."""
    persona = G.STRICTNESS_PRESETS[strictness]
    parts = [
        f"{persona} You are grading Question {q} of the KAUST Stage 3 Practical AI exam.",
        "",
        "## EXAM GUIDELINES (provided to students)",
        guidelines,
        "",
        rubric_item_block(q, row, idx, n_rows),
        "",
    ]
    if solution_code is not None:
        parts += ["## REFERENCE SOLUTION (instructor's notebook)", "```python", solution_code, "```", ""]
    parts += [
        "## STUDENT SUBMISSION (the one you must grade)",
        "```python",
        student_code,
        "```",
        "",
        "## YOUR TASK",
    ]
    if row["kind"] == "penalty":
        parts.append(
            f"Decide ONLY whether the penalty above applies. Return `awarded` = {row['points']:g} "
            f"if it applies and 0 if it does not."
        )
    else:
        parts.append(
            f"Grade ONLY the rubric item above, strictly against it. Award between 0 and "
            f"{row['points']:g} points for this item."
        )
    parts += [
        "Return a JSON object with: `awarded` (number) and `reasoning` (one short paragraph, "
        "<=120 words, explaining the award for this item only).",
        "OUTPUT RULES (strict):",
        "  - Output ONLY the JSON object. No prose before or after.",
        "  - Do NOT grade or mention other rubric items.",
        "  - Do NOT recalculate or revise the award inside `reasoning`.",
        "GRADING SCALE (strict):",
        "  - 100% of a task: code is correct and runs.",
        "  - 25-75% of a task: code is attempted with syntax/logic mistakes.",
        "  - 0% of a task: NO CODE was written for that task. This is the ONLY case for 0.",
        "  - -50% of a task: code uses an entirely wrong approach OR is blind copy-paste "
        "    OR disregards explicit instructions (e.g. wrong dataset).",
        "IMPORTANT: A submission that is disorganized, fragmented, uses a different "
        "structure than the reference, or fails to run end-to-end is NOT automatically 0. "
        "If the student wrote code that attempts this item, award partial credit even when "
        "other parts are broken. Only assign 0 when there is literally no code for it.",
    ]
    return "\n".join(parts)


def build_arbiter_prompt(q: int, rubric_section: str, guidelines: str, student_code: str,
                         solution_code: Optional[str], grade_a: dict, grade_b: dict) -> str:
    """Neutral-persona arbiter over the strict (A) and neutral (B) grades."""
    max_score, max_bonus, _ = G.QUESTION_META[q]
    persona = G.STRICTNESS_PRESETS["neutral"]
    parts = [
        f"{persona} You are the ARBITER for Question {q} of the KAUST Stage 3 Practical AI exam.",
        "Two graders assessed the same submission against the same rubric and disagree. "
        "Your job is to produce the final grade by resolving every disagreement with the "
        "rubric's partial-credit scale below, task by task, using the code itself as evidence. "
        "Neither grader is authoritative; a grader who withheld all credit for a task that "
        "contains an attempt violates the grading scale, and a grader who awarded full credit "
        "for code that does not meet the rubric row violates it too.",
        "",
        "## EXAM GUIDELINES (provided to students)",
        guidelines,
        "",
        f"## RUBRIC FOR QUESTION {q}",
        rubric_section,
        "",
    ]
    if solution_code is not None:
        parts += ["## REFERENCE SOLUTION (instructor's notebook)", "```python", solution_code, "```", ""]
    parts += [
        "## STUDENT SUBMISSION (the one being graded)",
        "```python",
        student_code,
        "```",
        "",
        "## GRADER A (instructed to be harsh)",
        "```json",
        json.dumps(grade_a, ensure_ascii=False, indent=1),
        "```",
        "",
        "## GRADER B (instructed to be strict but fair)",
        "```json",
        json.dumps(grade_b, ensure_ascii=False, indent=1),
        "```",
        "",
        "## YOUR TASK",
        f"Produce the final grade. Max base score for Q{q} is {max_score:g}; max bonus is {max_bonus:g}.",
        "Return a JSON object with: `score` (number, <= max base), `bonus` (number, <= max bonus; "
        "0 if no bonus tasks), `task_breakdown` (array of per-task awards, ONE entry for EVERY "
        "task row in the rubric above, including bonus rows), `resolutions` (array of short "
        "strings, one per task where A and B disagreed, naming which award you kept and why), "
        "and `reasoning` (one short paragraph, <=200 words).",
        "Sum the `awarded` values to get `score` (base tasks) and `bonus` (bonus tasks). Decide "
        "each task ONCE; do not revise.",
        "OUTPUT RULES (strict):",
        "  - Output ONLY the JSON object. No prose before or after.",
        "  - Each task_breakdown entry appears EXACTLY ONCE. Never repeat a task.",
        "  - Do NOT recalculate, re-sum, or revise scores inside `reasoning`.",
        "GRADING SCALE (strict):",
        "  - 100% of a task: code is correct and runs.",
        "  - 25-75% of a task: code is attempted with syntax/logic mistakes.",
        "  - 0% of a task: NO CODE was written for that task. This is the ONLY case for 0.",
        "  - -50% of a task: code uses an entirely wrong approach OR is blind copy-paste "
        "    OR disregards explicit instructions (e.g. wrong dataset).",
        "IMPORTANT: A submission that is disorganized, fragmented, uses a different "
        "structure than the reference, or fails to run end-to-end is NOT automatically 0. "
        "Grade each task individually: if the student wrote code that attempts the task, "
        "award partial credit even when other parts are broken. Only assign 0 to a task "
        "when there is literally no code for it.",
    ]
    return "\n".join(parts)


# ---------------------------------------------------------------------------
# model calls
# ---------------------------------------------------------------------------

class Backend:
    def __init__(self, a):
        self.a = a
        self.no_chat_template_kwargs = False
        if a.provider == "openai":
            import closed_batch_grade as C  # noqa: WPS433
            self.C = C
            self.client = C.make_client("openai")
        else:
            self.client = G.get_client(a.api_base, "EMPTY")

    def complete(self, prompt: str) -> tuple[str, str]:
        """Returns (text, finish_reason)."""
        a = self.a
        if a.provider == "openai":
            body = self.C.openai_body(prompt, a.model, a.max_output_tokens, a.reasoning, "object", 0.0)
            resp = self.client.chat.completions.create(**body)
        else:
            # Same knobs as the paper's grader (top_k -1, thinking off); no guided_json,
            # because the decomposed/arbiter replies use their own JSON shapes and the
            # paper's stack (vLLM 0.19.1) ignored the schema anyway (Appendix A).
            extra_body = {"top_k": -1}
            if not self.no_chat_template_kwargs:
                extra_body["chat_template_kwargs"] = {"enable_thinking": False}
            try:
                resp = self.client.chat.completions.create(
                    model=a.model, messages=[{"role": "user", "content": prompt}],
                    temperature=0.0, top_p=1.0, max_tokens=a.max_output_tokens, extra_body=extra_body)
            except Exception as e:  # noqa: BLE001
                # vLLM's Mistral tokenizer rejects chat_template_kwargs outright ("chat_template
                # is not supported for Mistral tokenizers"); the paper's ML grader drops the knob
                # after the first rejection, and so do we (a no-op for models without thinking).
                if "chat_template" in str(e) and not self.no_chat_template_kwargs:
                    logging.warning("server rejects chat_template_kwargs; dropping it for this run: %s", str(e)[:120])
                    self.no_chat_template_kwargs = True
                    extra_body.pop("chat_template_kwargs", None)
                    resp = self.client.chat.completions.create(
                        model=a.model, messages=[{"role": "user", "content": prompt}],
                        temperature=0.0, top_p=1.0, max_tokens=a.max_output_tokens, extra_body=extra_body)
                else:
                    raise
        ch = resp.choices[0]
        return ch.message.content or "", ch.finish_reason or ""


def call_json(be: Backend, prompt: str, log_rec: dict, calls_path: Path) -> Optional[dict]:
    """One call with the paper grader's retry policy; appends an audit record."""
    last: Optional[Exception] = None
    parsed: Optional[dict] = None
    text, finish, t0 = "", "", time.time()
    for attempt in range(1, 4):
        try:
            text, finish = be.complete(prompt)
            if finish == "length":
                raise RuntimeError(f"response truncated at max_tokens ({be.a.max_output_tokens})")
            parsed = G._parse_json_lenient(text)
            break
        except Exception as e:  # noqa: BLE001
            last = e
            if "truncated" in str(e) or isinstance(e, json.JSONDecodeError):
                logging.error("deterministic failure on attempt %d, not retrying: %s", attempt, e)
                break
            logging.warning("call failed (attempt %d/3): %s", attempt, e)
            time.sleep(2 * attempt)
    rec = dict(log_rec)
    rec.update({"prompt_sha1": hashlib.sha1(prompt.encode("utf-8")).hexdigest(),
                "prompt_chars": len(prompt), "finish": finish, "latency_s": round(time.time() - t0, 2),
                "raw": text, "parsed": parsed, "error": None if parsed is not None else str(last)})
    with LOCK:
        calls_path.parent.mkdir(parents=True, exist_ok=True)
        with open(calls_path, "a", encoding="utf-8") as f:
            f.write(json.dumps(rec, ensure_ascii=False) + "\n")
    return parsed


# ---------------------------------------------------------------------------
# grading modes
# ---------------------------------------------------------------------------

def _clamp(v: float, lo: float, hi: float) -> float:
    return max(lo, min(v, hi))


def grade_decomposed(be: Backend, a, q: int, rubric: str, guidelines: str, student_code: str,
                     solution_code: Optional[str], student: int, calls_path: Path):
    rows = parse_rubric_rows(rubric, q)
    if a.max_rows:
        rows = rows[: a.max_rows]
    n = len(rows)
    max_score, max_bonus, _ = G.QUESTION_META[q]

    def one(i_row):
        i, row = i_row
        prompt = build_decomposed_prompt(q, row, i + 1, n, guidelines, student_code, solution_code, a.persona)
        parsed = call_json(be, prompt, {"student": student, "q": q, "mode": "decomposed",
                                        "row": i + 1, "task": row["task"], "kind": row["kind"],
                                        "points": row["points"]}, calls_path)
        if parsed is None:
            return i, None
        try:
            v = float(parsed.get("awarded", 0) or 0)
        except (TypeError, ValueError):
            v = None
        return i, (v, parsed.get("reasoning", ""))

    with cf.ThreadPoolExecutor(max_workers=a.concurrency) as ex:
        out = dict(ex.map(one, list(enumerate(rows))))
    failed = [i for i, v in out.items() if v is None or v[0] is None]
    if failed:
        return None  # a missing row makes the question total undefined; leave the cell blank
    base = bonus = 0.0
    lines = []
    for i, row in enumerate(rows):
        v, why = out[i]
        if row["kind"] == "penalty":
            v = _clamp(v, row["points"], 0.0)
            base += v
        elif row["kind"] == "bonus":
            v = _clamp(v, 0.0, row["points"])
            bonus += v
        else:
            v = _clamp(v, 0.0, row["points"])
            base += v
        lines.append(f"[{row['kind']}] {row['task']}: {v:g}/{row['points']:g} -- {str(why).strip()[:160]}")
    score = _clamp(base, 0.0, max_score)
    bonus = _clamp(bonus, 0.0, max_bonus)
    reasoning = f"DECOMPOSED ({n} rubric items, {a.persona} persona; raw base sum {base:g}):\n" + "\n".join(lines)
    return score, bonus, reasoning


def grade_arbitrated(be: Backend, a, q: int, rubric: str, guidelines: str, student_code: str,
                     solution_code: Optional[str], student: int, calls_path: Path):
    def paper_grade(strictness: str, stage: str):
        prompt = G.build_prompt(q=q, rubric_section=rubric, guidelines=guidelines, student_code=student_code,
                                solution_code=solution_code, use_breakdown=True, use_guidelines=True,
                                strictness=strictness, few_shot_examples=None)
        return call_json(be, prompt, {"student": student, "q": q, "mode": "arbitrated", "stage": stage,
                                      "persona": strictness}, calls_path)

    with cf.ThreadPoolExecutor(max_workers=min(2, a.concurrency)) as ex:
        fa = ex.submit(paper_grade, a.persona, "A_strict")
        fb = ex.submit(paper_grade, "neutral", "B_neutral")
        ga, gb = fa.result(), fb.result()
    if ga is None or gb is None:
        return None
    prompt = build_arbiter_prompt(q, rubric, guidelines, student_code, solution_code, ga, gb)
    gc = call_json(be, prompt, {"student": student, "q": q, "mode": "arbitrated", "stage": "C_arbiter"}, calls_path)
    if gc is None:
        return None
    max_score, max_bonus, _ = G.QUESTION_META[q]
    try:
        score = _clamp(float(gc.get("score", 0) or 0), 0.0, max_score)
        bonus = _clamp(float(gc.get("bonus", 0) or 0), 0.0, max_bonus)
    except (TypeError, ValueError):
        return None

    def s(g):
        try:
            return f"{float(g.get('score', 0) or 0):g}/{float(g.get('bonus', 0) or 0):g}"
        except (TypeError, ValueError):
            return "?"
    reasoning = (f"ARBITRATED: A(strict) {s(ga)}  B(neutral) {s(gb)}  -> arbiter {score:g}/{bonus:g}. "
                 f"Resolutions: {json.dumps(gc.get('resolutions', []), ensure_ascii=False)[:600]} "
                 f"Reasoning: {str(gc.get('reasoning', '')).strip()[:600]}")
    return score, bonus, reasoning


# ---------------------------------------------------------------------------
# driver
# ---------------------------------------------------------------------------

def parse_students(spec: str) -> list[int]:
    out: list[int] = []
    for part in spec.split(","):
        part = part.strip()
        if "-" in part:
            lo, hi = part.split("-")
            out += list(range(int(lo), int(hi) + 1))
        elif part:
            out.append(int(part))
    return out


def cmd_grade(a) -> None:
    RESULTS.mkdir(parents=True, exist_ok=True)
    CALLS.mkdir(parents=True, exist_ok=True)
    tag = (f"{a.run_id}_m-{a.model_tag}_mode-{a.mode}_sol1_gd1_rs0_bd1_str-{a.persona}_t00_n1"
           f"_s{a.students.replace(',', '_')}")
    out_xlsx = RESULTS / f"{a.run_id}__{tag}.xlsx"
    calls_path = CALLS / f"{a.run_id}.jsonl"
    G.setup_logging(OUT_ROOT / f"{a.run_id}.log")
    logging.info("run %s: mode=%s model=%s backend=%s students=%s questions=%s -> %s",
                 a.run_id, a.mode, a.model, a.provider or a.api_base, a.students, a.questions, out_xlsx)

    if not out_xlsx.exists():
        frozen = G.init_output_workbook(G.GRADES_XLSX, out_xlsx)
        logging.info("cloned grades workbook (%d formula cells frozen)", frozen)
    wb = G.safe_load_workbook(out_xlsx, G.GRADES_XLSX)
    ws = wb.active
    cols = G.ensure_ai_columns(ws)
    row_index = G.build_student_row_index(ws)

    rubrics = G.load_rubric_sections()
    guidelines = G.GUIDELINES.read_text(encoding="utf-8")
    be = Backend(a)
    questions = [int(x) for x in a.questions.split(",")]

    for student in parse_students(a.students):
        row = row_index.get(student)
        sdir = G.EXTRACTED_DIR / str(student)
        if row is None or not sdir.is_dir():
            logging.warning("student %s: no row / no extracted folder; skipped", student)
            continue
        total_score = total_bonus = 0.0
        complete = True
        for q in questions:
            score_col, reason_col = cols[f"AI Q{q} Score"], cols[f"AI Q{q} Reasoning"]
            bonus_col = cols.get(f"AI Q{q} Bonus")
            existing = ws.cell(row=row, column=score_col).value
            if existing not in (None, ""):
                total_score += float(existing)
                if bonus_col and ws.cell(row=row, column=bonus_col).value not in (None, ""):
                    total_bonus += float(ws.cell(row=row, column=bonus_col).value)
                continue
            nb = sdir / f"Q{q}.ipynb"
            if not nb.exists():
                ws.cell(row=row, column=score_col, value=0)
                if bonus_col:
                    ws.cell(row=row, column=bonus_col, value=0)
                ws.cell(row=row, column=reason_col, value="No submission found.")
                G.atomic_save(wb, out_xlsx)
                continue
            student_code = G.notebook_to_text(nb)
            sol = G.SOLUTIONS_DIR / G.QUESTION_META[q][2]
            solution_code = G.notebook_to_text(sol) if sol.exists() else None
            t0 = time.time()
            fn = grade_decomposed if a.mode == "decomposed" else grade_arbitrated
            res = fn(be, a, q, rubrics[q], guidelines, student_code, solution_code, student, calls_path)
            if res is None:
                logging.error("student %s Q%d: grading failed; cell left blank", student, q)
                complete = False
                continue
            score, bonus, reasoning = res
            ws.cell(row=row, column=score_col, value=score)
            if bonus_col:
                ws.cell(row=row, column=bonus_col, value=bonus)
            ws.cell(row=row, column=reason_col, value=reasoning[:32000])
            total_score += score
            total_bonus += bonus
            G.atomic_save(wb, out_xlsx)
            logging.info("student %s Q%d: score=%g bonus=%g (%.0fs)", student, q, score, bonus, time.time() - t0)
        if complete:
            ws.cell(row=row, column=cols["AI Total Score"], value=total_score)
            ws.cell(row=row, column=cols["AI Total Bonus"], value=total_bonus)
            G.atomic_save(wb, out_xlsx)
            logging.info("student %s: total=%g bonus=%g", student, total_score, total_bonus)
    logging.info("Done. Workbook: %s", out_xlsx)


# ---------------------------------------------------------------------------
# report
# ---------------------------------------------------------------------------

def _load_totals(xlsx: Path) -> dict[int, dict]:
    """student -> {ai, ta, q: {k: (ai, ta)}} for rows with a numeric AI total (Q1-Q3 base)."""
    ws = openpyxl.load_workbook(xlsx, data_only=True).active
    hdr = {c.value: i for i, c in enumerate(ws[1]) if c.value}
    out: dict[int, dict] = {}

    def f(v):
        try:
            return float(v)
        except (TypeError, ValueError):
            return None
    for r in ws.iter_rows(min_row=2, values_only=True):
        if r[0] is None:
            continue
        sid = int(r[0])
        qs = {}
        ok = True
        for k in (1, 2, 3):
            ai = f(r[hdr[f"AI Q{k} Score"]])
            t1, t2 = f(r[hdr[f"TA 1 - Q{k} Score"]]), f(r[hdr[f"TA 2 - Q{k} Score"]])
            if ai is None or t1 is None or t2 is None:
                ok = False
                break
            qs[k] = (ai, (t1 + t2) / 2)
        if not ok:
            continue
        ai_tot = sum(v[0] for v in qs.values())
        ta_tot = sum(v[1] for v in qs.values())
        out[sid] = {"ai": ai_tot, "ta": ta_tot, "q": qs}
    return out


def _stats(d: dict[int, dict], students: list[int]) -> dict:
    xs = [d[s] for s in students if s in d]
    n = len(xs)
    if n == 0:
        return {"n": 0}
    res = [x["ai"] - x["ta"] for x in xs]
    tot = [x["ai"] for x in xs]
    mae = statistics.fmean(abs(e) for e in res)
    sd = statistics.pstdev(tot) if n > 1 else 0.0
    zero = sum(1 for t in tot if t == 0) / n
    beh = ("refusal" if zero >= 0.90 and sd < 0.5 else "near-refusal" if zero >= 0.90
           else "collapse" if mae >= 8 else "graded")
    pq = {k: statistics.fmean(abs(x["q"][k][0] - x["q"][k][1]) for x in xs) for k in (1, 2, 3)}
    return {"n": n, "mae": mae, "bias": statistics.fmean(res), "zero_rate": zero, "sd": sd,
            "mean_total": statistics.fmean(tot), "behaviour": beh, "per_q_mae": pq}


def _find_paper_run(rid: str) -> Path:
    hits = sorted((G.RESULTS_DIR).glob(f"{rid}__*.xlsx"))
    if not hits:
        sys.exit(f"no paper workbook for {rid} under {G.RESULTS_DIR}")
    return hits[0]


def cmd_report(a) -> None:
    hits = sorted(RESULTS.glob(f"{a.run_id}__*.xlsx"))
    if not hits:
        sys.exit(f"no workbook for {a.run_id} under {RESULTS}")
    mine = _load_totals(hits[0])
    students = sorted(mine)
    print(f"# {a.run_id}: {hits[0].name}")
    print(f"graded students with complete Q1-Q3: {len(students)}  (range {students[0] if students else '-'}..{students[-1] if students else '-'})")
    rows = [("this run", _stats(mine, students))]
    for label, rid in (("paper strict", a.paper_strict), ("paper neutral", a.paper_neutral)):
        if rid:
            rows.append((f"{label} {rid}", _stats(_load_totals(_find_paper_run(rid)), students)))
    print(f"{'run':<28}{'n':>5}{'MAE':>8}{'bias':>8}{'zero%':>7}{'sd':>7}{'mean':>7}  behaviour   Q1/Q2/Q3 MAE")
    for label, s in rows:
        if s["n"] == 0:
            print(f"{label:<28}{0:>5}  (no overlapping students)")
            continue
        pq = "/".join(f"{s['per_q_mae'][k]:.2f}" for k in (1, 2, 3))
        print(f"{label:<28}{s['n']:>5}{s['mae']:>8.3f}{s['bias']:>+8.3f}{100*s['zero_rate']:>6.1f}%{s['sd']:>7.2f}"
              f"{s['mean_total']:>7.2f}  {s['behaviour']:<11} {pq}")


def main() -> None:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    sub = ap.add_subparsers(dest="cmd", required=True)
    g = sub.add_parser("grade")
    g.add_argument("--mode", choices=["decomposed", "arbitrated"], required=True)
    g.add_argument("--run-id", required=True)
    g.add_argument("--model", required=True)
    g.add_argument("--model-tag", default=None)
    g.add_argument("--api-base", default=None)
    g.add_argument("--provider", choices=["openai"], default=None)
    g.add_argument("--reasoning", default="none", help="OpenAI reasoning_effort (provider=openai only)")
    g.add_argument("--students", default="1-100")
    g.add_argument("--questions", default="1,2,3")
    g.add_argument("--persona", default="strict", choices=sorted(G.STRICTNESS_PRESETS))
    g.add_argument("--max-output-tokens", type=int, default=4096)
    g.add_argument("--concurrency", type=int, default=4)
    g.add_argument("--max-rows", type=int, default=0, help="smoke only: limit rubric rows per question")
    g.set_defaults(func=cmd_grade)
    r = sub.add_parser("report")
    r.add_argument("--run-id", required=True)
    r.add_argument("--paper-strict", default=None, help="paper run id to compare, e.g. L-C01")
    r.add_argument("--paper-neutral", default=None, help="paper run id to compare, e.g. L-A32")
    r.set_defaults(func=cmd_report)
    a = ap.parse_args()
    if a.cmd == "grade":
        if not a.provider and not a.api_base:
            ap.error("give --api-base (vLLM) or --provider openai")
        if a.provider and a.api_base:
            ap.error("give only one backend")
        if not a.model_tag:
            a.model_tag = re.sub(r"[^A-Za-z0-9.]+", "-", a.model.split("/")[-1]).lower()
    a.func(a)


if __name__ == "__main__":
    main()
